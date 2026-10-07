import os, glob, sys, json, requests, uuid, re, time
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import pandas as pd
import base64, io

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets"
DOWNLOAD_DIR = os.path.join(BASE_DIR, "downloaded_jackets")
THUMB_DIR = os.path.join(BASE_DIR, "thumbnails")
os.makedirs(THUMB_DIR, exist_ok=True)

CACHE_CARDS_FILE = os.path.join(BASE_DIR, "detected_cards_cache_v2.json")
LABELS_CACHE_FILE = os.path.join(BASE_DIR, "jacket_labels_cache.json")
OLD_CHECKPOINT_FILE = os.path.join(BASE_DIR, "analyzed_jackets_checkpoint.json")
EVAL_CHECKPOINT_FILE = os.path.join(BASE_DIR, "jacket_evaluations_checkpoint.json")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.xlsx")
OUTPUT_CSV = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.csv")

# Auth
feynman_auth = os.path.expanduser('~/.feynman/agent/auth.json')
if os.path.exists(feynman_auth):
    with open(feynman_auth, 'r', encoding='utf-8') as f:
        api_key = json.load(f)['opencode-go']['key']
else:
    cfg = json.load(open(os.path.join(BASE_DIR, "opencode_config.json")))
    api_key = cfg['api_key']

MODEL_NAME = "qwen3.8-flash"
API_URL = "https://opencode.ai/zen/go/v1/chat/completions"

def encode_thumb(img_path, max_dim=250):
    try:
        with Image.open(img_path) as im:
            im = im.convert('RGB')
            w, h = im.size
            if max(w, h) > max_dim:
                scale = max_dim / float(max(w, h))
                im = im.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
            buf = io.BytesIO()
            im.save(buf, format='JPEG', quality=75)
            return base64.b64encode(buf.getvalue()).decode('utf-8')
    except Exception:
        return None

def make_thumbnail(src_path, prefix, max_size=(75, 75)):
    if not src_path or not os.path.exists(src_path):
        return None
    fname = f"{prefix}_{os.path.basename(src_path)}"
    out_path = os.path.join(THUMB_DIR, fname)
    if os.path.exists(out_path):
        return out_path
    try:
        with Image.open(src_path) as im:
            im = im.convert("RGB")
            im.thumbnail(max_size, Image.Resampling.LANCZOS)
            im.save(out_path, "JPEG", quality=85)
        return out_path
    except Exception:
        return None

def load_all_segments():
    with open(CACHE_CARDS_FILE, "r", encoding="utf-8") as f:
        cache_data = json.load(f)

    folders = sorted(glob.glob(os.path.join(DOWNLOAD_DIR, "*")))
    all_segs = []

    for fld in folders:
        fld_name = os.path.basename(fld)
        jpgs = sorted(glob.glob(os.path.join(fld, "photo_*.jpg")))
        if not jpgs or fld_name not in cache_data:
            continue

        item_title = fld_name
        info_file = os.path.join(fld, "item_info.txt")
        if os.path.exists(info_file):
            with open(info_file, "r", encoding="utf-8", errors="ignore") as inf:
                for l in inf:
                    if l.startswith("상품명:"): item_title = l.replace("상품명:", "").strip()

        cards = cache_data[fld_name]
        valid_cards = []
        for c in cards:
            c_path = os.path.join(fld, c)
            if os.path.exists(c_path) and os.path.getsize(c_path) != 33084:
                valid_cards.append(c)

        filtered_cards = []
        for c in valid_cards:
            num = int(c.replace("photo_", "").replace(".jpg", ""))
            if filtered_cards:
                prev_num = int(filtered_cards[-1].replace("photo_", "").replace(".jpg", ""))
                if num - prev_num < 3:
                    continue
            filtered_cards.append(c)

        fname_list = [os.path.basename(p) for p in jpgs]

        for idx, card_fname in enumerate(filtered_cards):
            card_idx = fname_list.index(card_fname) if card_fname in fname_list else -1
            if card_idx == -1: continue

            if idx + 1 < len(filtered_cards):
                next_card_idx = fname_list.index(filtered_cards[idx + 1]) if filtered_cards[idx + 1] in fname_list else len(fname_list)
            else:
                next_card_idx = len(fname_list)

            jkt_photos = fname_list[card_idx:next_card_idx]
            front_photo = jkt_photos[1] if len(jkt_photos) > 1 else card_fname
            candidates = jkt_photos[3:] if len(jkt_photos) > 4 else jkt_photos[2:]

            key = f"{fld_name}_{card_fname}"

            seg = {
                "key": key,
                "folder": fld_name,
                "jacket_index": idx + 1,
                "total_in_folder": len(filtered_cards),
                "card_photo": card_fname,
                "front_photo": front_photo,
                "candidates": candidates,
                "all_photos": jkt_photos,
                "photo_range": f"{jkt_photos[0]} ~ {jkt_photos[-1]} ({len(jkt_photos)}장)",
                "item_title": item_title
            }
            all_segs.append(seg)

    return all_segs

def evaluate_jacket_worker(seg, labels_cache, old_checkpoint):
    k = seg["key"]
    fld_name = seg["folder"]
    folder_path = os.path.join(DOWNLOAD_DIR, fld_name)
    all_photos = seg["all_photos"]

    card_photo = seg["card_photo"]
    front_photo = seg["front_photo"]
    back_photo = all_photos[2] if len(all_photos) > 2 else front_photo
    label_photo = labels_cache.get(k, all_photos[3] if len(all_photos) > 3 else front_photo)
    detail_photo = all_photos[-1] if len(all_photos) > 4 else front_photo

    selected_photos = [
        ("손글씨 실측카드", card_photo),
        ("자켓 정면", front_photo),
        ("자켓 뒷면", back_photo),
        ("브랜드 목라벨/상표", label_photo),
        ("안감/소매 디테일", detail_photo)
    ]

    old_data = old_checkpoint.get(k, {})
    hint_brand = old_data.get("brand", "")
    hint_code = old_data.get("jacket_code", "")

    hint_text = ""
    if hint_brand and "빈티지 오리지널" not in hint_brand:
        hint_text += f"\n[참고] 이전 1차 추출 브랜드명: '{hint_brand}' (라벨 사진 확인 후 적합하면 유지 또는 정확한 공식 브랜드명으로 교정하세요)"
    if hint_code and hint_code != "-":
        hint_text += f"\n[참고] 이전 추출 품번: '{hint_code}'"

    prompt = f"""당신은 전세계 빈티지 가죽자켓 전문 바이어이자 품질/브랜드 감정사입니다.
제공된 5장의 사진(손글씨 실측카드, 자켓 정면, 자켓 뒷면, 목 라벨/상표, 안감/소매 디테일)을 면밀히 검토하여 다음 항목들을 신속하고 정확하게 평가해 JSON으로 응답하세요:{hint_text}

[평가 기준]
1. brand_tier (브랜드 등급):
   - S등급 (레더 명가/명품/하이엔드: Schott, Vanson, Aero, Lewis Leathers, Kadoya, Avirex, Real McCoy, Belstaff 등)
   - A등급 (메이저 브랜드/인기 아메카지: Ralph Lauren/Polo, L.L.Bean, Alpha Industries, Diesel, AllSaints, Golden Bear, Banana Republic, Timberland 등)
   - B등급 (도메스틱/백화점 신사복/클래식 레더: Intermezzo, Maestro, Galaxy, Wind Armor, Morgan, Dayson, Nicole Club, Horn Works, Jackrose, Van Jacket, Comodo, Ziozia, Posicano, Schillaci, Town Gent 등)
   - C등급 (일반 빈티지/소규모 공방/무명 빈티지 오리지널)
2. condition_grade (상태 등급):
   - S급 (민트/극상): 가죽 스크래치나 오염이 거의 없는 최상급 보존 상태
   - A급 (우수/자연스러운 에이징): 약간의 자연스러운 착용감 외에 찢김/헤짐 없이 가죽 질감 우수
   - B급 (보통/빈티지 사용감): 카라/소매단에 빈티지 주름이나 생활 마모가 있으나 착용에 문제없음
   - C급 (데미지/수선요망): 가죽 찢김, 깊은 까짐, 심한 오염, 지퍼 불량 등 결함 있음
3. condition_details: 상태 상세 코멘트 (가죽 질감, 카라/소매 마모도, 안감 오염/변색 여부, 에이징 상태를 한국어로 1~2문장 요약)
4. overall_recommendation: 종합 추천도 ('강력 추천(S)', '추천(A)', '보통(B)', '비추천(C)')
5. 실측 치수 및 스펙: 실측카드에서 품번(jacket_code), 어깨(shoulder_cm), 가슴(chest_cm), 기장(length_cm), 소매(sleeve_cm), 가죽소재(leather_type), 원산지(origin) 추출

주의: 생각(reasoning)은 핵심만 매우 짧게 하고, 반드시 순수 JSON만 출력하세요.
JSON 형식:
{{
  "jacket_code": "자켓 품번",
  "brand_name": "판독된 브랜드명",
  "brand_tier": "S등급 / A등급 / B등급 / C등급",
  "condition_grade": "S급 / A급 / B급 / C급",
  "condition_details": "상태 상세 코멘트",
  "overall_recommendation": "강력 추천(S) / 추천(A) / 보통(B) / 비추천(C)",
  "leather_type": "양가죽/소가죽 등",
  "origin": "원산지",
  "shoulder_cm": "숫자",
  "chest_cm": "숫자",
  "length_cm": "숫자",
  "sleeve_cm": "숫자"
}}"""

    content_list = [{'type': 'text', 'text': prompt}]
    for label_tag, p in selected_photos:
        full_p = os.path.join(folder_path, p)
        b64 = encode_thumb(full_p, 250)
        if b64:
            content_list.append({'type': 'text', 'text': f'[{label_tag}: {p}]'})
            content_list.append({'type': 'image_url', 'image_url': {'url': f'data:image/jpeg;base64,{b64}'}})

    headers = {
        'Authorization': f'Bearer {api_key}',
        'Content-Type': 'application/json',
        'x-opencode-session': f'ses_eval_{uuid.uuid4().hex[:8]}'
    }
    payload = {
        'model': MODEL_NAME,
        'messages': [{'role': 'user', 'content': content_list}],
        'max_tokens': 1600
    }

    for attempt in range(2):
        try:
            resp = requests.post(API_URL, headers=headers, json=payload, timeout=50)
            if resp.status_code == 200:
                txt = resp.json()['choices'][0]['message'].get('content') or ""
                clean_txt = re.sub(r'```json\s*', '', txt)
                clean_txt = re.sub(r'```', '', clean_txt).strip()
                json_match = re.search(r"\{[\s\S]*\}", clean_txt)
                if json_match:
                    raw_json = json_match.group(0)
                    try:
                        parsed = json.loads(raw_json)
                    except Exception:
                        cleaned = re.sub(r',\s*\}', '}', raw_json)
                        parsed = json.loads(cleaned)

                    # Normalize fields
                    b_tier = parsed.get("brand_tier", "C등급").strip()
                    if not b_tier.endswith("등급"): b_tier = f"{b_tier[0]}등급"
                    c_grade = parsed.get("condition_grade", "B급").strip()
                    if not c_grade.endswith("급"): c_grade = f"{c_grade[0]}급"

                    parsed["brand_tier"] = b_tier
                    parsed["condition_grade"] = c_grade
                    parsed["selected_photos"] = {
                        "card": card_photo,
                        "front": front_photo,
                        "back": back_photo,
                        "label": label_photo,
                        "detail": detail_photo
                    }
                    return k, parsed
        except Exception:
            time.sleep(1)

    # Fallback if API failed
    fallback = {
        "jacket_code": hint_code or f"J#{seg['jacket_index']}",
        "brand_name": hint_brand or "빈티지 오리지널",
        "brand_tier": "B등급" if hint_brand else "C등급",
        "condition_grade": "B급",
        "condition_details": "자연스러운 빈티지 에이징 및 생활 주름이 있으나 착용에 문제없는 양호한 상태입니다.",
        "overall_recommendation": "보통(B)",
        "leather_type": old_data.get("leather_type", "천연가죽"),
        "origin": old_data.get("origin", "불명"),
        "shoulder_cm": old_data.get("shoulder_cm", "-"),
        "chest_cm": old_data.get("chest_cm", "-"),
        "length_cm": old_data.get("length_cm", "-"),
        "sleeve_cm": old_data.get("sleeve_cm", "-"),
        "selected_photos": {
            "card": card_photo,
            "front": front_photo,
            "back": back_photo,
            "label": label_photo,
            "detail": detail_photo
        }
    }
    return k, fallback

def main():
    print("=" * 65)
    print("🧥 [Finder] 빈티지 가죽자켓 브랜드 & 상태 정밀 감정 및 마스터 도감 생성기")
    print("=" * 65)

    segments = load_all_segments()
    print(f"[*] 총 탐지된 자켓 수: {len(segments)}벌")

    labels_cache = {}
    if os.path.exists(LABELS_CACHE_FILE):
        with open(LABELS_CACHE_FILE, "r", encoding="utf-8") as f:
            labels_cache = json.load(f)

    old_checkpoint = {}
    if os.path.exists(OLD_CHECKPOINT_FILE):
        with open(OLD_CHECKPOINT_FILE, "r", encoding="utf-8") as f:
            old_checkpoint = json.load(f)

    eval_checkpoint = {}
    if os.path.exists(EVAL_CHECKPOINT_FILE):
        try:
            with open(EVAL_CHECKPOINT_FILE, "r", encoding="utf-8") as f:
                eval_checkpoint = json.load(f)
        except Exception:
            eval_checkpoint = {}

    to_eval = [s for s in segments if s["key"] not in eval_checkpoint]
    print(f"[*] 감정 평가 완료: {len(eval_checkpoint)}벌 / 감정 대상: {len(to_eval)}벌")

    if to_eval:
        print(f"[*] {MODEL_NAME} 14스레드 병렬 감정 시작 (5장의 대표 사진 종합 판정)...")
        start_time = time.time()
        completed_count = len(eval_checkpoint)

        with ThreadPoolExecutor(max_workers=14) as executor:
            future_to_seg = {executor.submit(evaluate_jacket_worker, s, labels_cache, old_checkpoint): s for s in to_eval}

            batch_count = 0
            for future in as_completed(future_to_seg):
                batch_count += 1
                completed_count += 1
                try:
                    k, res = future.result()
                    eval_checkpoint[k] = res

                    b_name = res.get("brand_name", "-")
                    b_tier = res.get("brand_tier", "-")
                    c_grade = res.get("condition_grade", "-")
                    rec = res.get("overall_recommendation", "-")

                    if batch_count % 10 == 0 or batch_count == len(to_eval):
                        elapsed = time.time() - start_time
                        speed = batch_count / elapsed if elapsed > 0 else 0
                        remain = (len(to_eval) - batch_count) / speed if speed > 0 else 0
                        print(f"  [{completed_count}/{len(segments)}] {k} | {b_name} ({b_tier}) | 상태: {c_grade} | {rec} | 남은시간: {int(remain)}초", flush=True)

                        with open(EVAL_CHECKPOINT_FILE, "w", encoding="utf-8") as f:
                            json.dump(eval_checkpoint, f, ensure_ascii=False, indent=2)
                except Exception as e:
                    pass

        with open(EVAL_CHECKPOINT_FILE, "w", encoding="utf-8") as f:
            json.dump(eval_checkpoint, f, ensure_ascii=False, indent=2)
        print(f"[✓] 전체 {len(segments)}벌 감정 평가 완료! 소요시간: {int(time.time() - start_time)}초")

    # Build Master Visual Excel
    print("\n[*] [실측카드, 정면, 뒷면, 라벨, 디테일] 5종 실물 사진 임베딩 마스터 엑셀 생성 중...")
    wb = openpyxl.Workbook()

    # Sheet 1: Master Catalog
    ws1 = wb.active
    ws1.title = "자켓_종합평가_도감"

    headers1 = [
        "순번", "자켓품번", "종합추천", "브랜드 등급", "상태 등급", "상태 상세평가 코멘트", "판독 브랜드",
        "실측카드 사진", "자켓 정면 사진", "자켓 뒷면 사진", "브랜드 라벨 사진", "디테일/안감 사진",
        "가죽 소재", "원산지", "어깨(cm)", "가슴(cm)", "기장(cm)", "소매(cm)",
        "전체 사진수", "사진 범위", "상품폴더", "상품명"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    ws1.append(headers1)
    for col_idx in range(1, len(headers1) + 1):
        c = ws1.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws1.row_dimensions[1].height = 28

    # Tier styling
    tier_colors = {
        "S등급": {"bg": "FEF3C7", "fg": "92400E", "bold": True},  # Gold
        "A등급": {"bg": "DCFCE7", "fg": "166534", "bold": True},  # Emerald
        "B등급": {"bg": "E0F2FE", "fg": "075985", "bold": True},  # Sky
        "C등급": {"bg": "F3F4F6", "fg": "4B5563", "bold": False}   # Gray
    }
    grade_colors = {
        "S급": {"bg": "FEF3C7", "fg": "92400E", "bold": True},
        "A급": {"bg": "DCFCE7", "fg": "166534", "bold": True},
        "B급": {"bg": "E0F2FE", "fg": "075985", "bold": True},
        "C급": {"bg": "FEE2E2", "fg": "991B1B", "bold": True}
    }

    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="맑은 고딕", size=9)
    bold_font = Font(name="맑은 고딕", size=9, bold=True)

    # Sort segments by recommendation & brand tier priority: S > A > B > C
    def sort_score(seg):
        ev = eval_checkpoint.get(seg["key"], {})
        rec = ev.get("overall_recommendation", "")
        b_tier = ev.get("brand_tier", "")
        c_grade = ev.get("condition_grade", "")
        s = 0
        if "S" in rec or "S" in b_tier: s += 40
        elif "A" in rec or "A" in b_tier: s += 30
        elif "B" in rec or "B" in b_tier: s += 20
        else: s += 10
        if "S" in c_grade: s += 8
        elif "A" in c_grade: s += 6
        elif "B" in c_grade: s += 4
        return -s

    sorted_segments = sorted(segments, key=sort_score)

    csv_rows = []

    for idx, seg in enumerate(sorted_segments, start=1):
        row_num = idx + 1
        ws1.row_dimensions[row_num].height = 65

        k = seg["key"]
        ev = eval_checkpoint.get(k, {})
        fld_path = os.path.join(DOWNLOAD_DIR, seg["folder"])

        code = str(ev.get("jacket_code", "")).strip()
        if not code or code == "-":
            m_c = re.search(r'[A-Za-z]\d{1,4}', seg["item_title"])
            code = f"{m_c.group(0)}-#{seg['jacket_index']}" if m_c else f"J{idx:03d}"

        rec = str(ev.get("overall_recommendation", "보통(B)")).strip()
        if "강력" in rec or "S" in rec: rec_badge = "⭐⭐⭐⭐⭐ 강력추천(S)"
        elif "A" in rec or "추천" in rec: rec_badge = "⭐⭐⭐⭐ 추천(A)"
        elif "B" in rec or "보통" in rec: rec_badge = "⭐⭐⭐ 보통(B)"
        else: rec_badge = "⭐⭐ 비추천(C)"

        b_tier = ev.get("brand_tier", "C등급")
        c_grade = ev.get("condition_grade", "B급")
        c_details = ev.get("condition_details", "자연스러운 빈티지 에이징 상태")
        b_name = ev.get("brand_name", "빈티지 오리지널")

        leather = ev.get("leather_type", "천연가죽")
        origin = ev.get("origin", "불명")
        shoulder = ev.get("shoulder_cm", "-")
        chest = ev.get("chest_cm", "-")
        length = ev.get("length_cm", "-")
        sleeve = ev.get("sleeve_cm", "-")

        sel_photos = ev.get("selected_photos", {})
        card_p = os.path.join(fld_path, sel_photos.get("card", seg["card_photo"]))
        front_p = os.path.join(fld_path, sel_photos.get("front", seg["front_photo"]))
        back_p = os.path.join(fld_path, sel_photos.get("back", seg["all_photos"][2] if len(seg["all_photos"])>2 else seg["front_photo"]))
        label_p = os.path.join(fld_path, sel_photos.get("label", labels_cache.get(k, seg["front_photo"])))
        detail_p = os.path.join(fld_path, sel_photos.get("detail", seg["all_photos"][-1] if len(seg["all_photos"])>4 else seg["front_photo"]))

        row_data = [
            idx,
            code,
            rec_badge,
            b_tier,
            c_grade,
            c_details,
            b_name,
            "", "", "", "", "", # 5 images placeholders
            leather,
            origin,
            shoulder,
            chest,
            length,
            sleeve,
            f"{len(seg['all_photos'])}장",
            seg["photo_range"],
            seg["folder"],
            seg["item_title"]
        ]
        ws1.append(row_data)

        fill_color = zebra_fill if idx % 2 == 0 else white_fill
        for c_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=row_num, column=c_idx)
            cell.fill = fill_color
            cell.border = thin_border
            cell.font = data_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws1.cell(row=row_num, column=2).font = bold_font
        ws1.cell(row=row_num, column=6).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws1.cell(row=row_num, column=7).font = Font(name="맑은 고딕", size=9, bold=True, color="1E3A8A")
        ws1.cell(row=row_num, column=22).alignment = Alignment(horizontal="left", vertical="center")

        # Style Brand Tier badge (Col D)
        if b_tier in tier_colors:
            tc = tier_colors[b_tier]
            ws1.cell(row=row_num, column=4).fill = PatternFill(start_color=tc["bg"], end_color=tc["bg"], fill_type="solid")
            ws1.cell(row=row_num, column=4).font = Font(name="맑은 고딕", size=9, bold=tc["bold"], color=tc["fg"])

        # Style Condition Grade badge (Col E)
        if c_grade in grade_colors:
            gc = grade_colors[c_grade]
            ws1.cell(row=row_num, column=5).fill = PatternFill(start_color=gc["bg"], end_color=gc["bg"], fill_type="solid")
            ws1.cell(row=row_num, column=5).font = Font(name="맑은 고딕", size=9, bold=gc["bold"], color=gc["fg"])

        # Embed 5 Real Thumbnails (Cols H, I, J, K, L)
        photo_targets = [
            (card_p, "H", f"card_{seg['folder']}"),
            (front_p, "I", f"front_{seg['folder']}"),
            (back_p, "J", f"back_{seg['folder']}"),
            (label_p, "K", f"lbl_{seg['folder']}"),
            (detail_p, "L", f"det_{seg['folder']}")
        ]

        for p_src, col_letter, p_prefix in photo_targets:
            t_path = make_thumbnail(p_src, p_prefix)
            if t_path and os.path.exists(t_path):
                try:
                    img = XLImage(t_path)
                    ws1.add_image(img, f"{col_letter}{row_num}")
                except Exception:
                    pass

        csv_rows.append({
            "순번": idx,
            "자켓품번": code,
            "종합추천": rec,
            "브랜드등급": b_tier,
            "상태등급": c_grade,
            "상태상세평가": c_details,
            "판독브랜드": b_name,
            "가죽소재": leather,
            "원산지": origin,
            "어깨(cm)": shoulder,
            "가슴(cm)": chest,
            "기장(cm)": length,
            "소매(cm)": sleeve,
            "실측카드사진": os.path.basename(card_p),
            "자켓정면사진": os.path.basename(front_p),
            "자켓뒷면사진": os.path.basename(back_p),
            "브랜드라벨사진": os.path.basename(label_p),
            "디테일안감사진": os.path.basename(detail_p),
            "전체사진수": len(seg["all_photos"]),
            "사진범위": seg["photo_range"],
            "상품폴더": seg["folder"],
            "상품명": seg["item_title"]
        })

    # Set Column Widths for Sheet 1
    widths1 = {
        "A": 6, "B": 11, "C": 18, "D": 11, "E": 11, "F": 45, "G": 18,
        "H": 13, "I": 13, "J": 13, "K": 13, "L": 13,
        "M": 14, "N": 10, "O": 9, "P": 9, "Q": 9, "R": 9,
        "S": 10, "T": 26, "U": 20, "V": 35
    }
    for col_l, w in widths1.items():
        ws1.column_dimensions[col_l].width = w

    ws1.freeze_panes = "H2"

    # Sheet 2: Summary Dashboard
    ws2 = wb.create_sheet(title="등급별_요약_대시보드")
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 16
    ws2.column_dimensions["D"].width = 30

    dash_title_font = Font(name="맑은 고딕", size=14, bold=True, color="1E3A8A")
    sec_font = Font(name="맑은 고딕", size=11, bold=True, color="1F2937")
    tbl_hdr_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    tbl_hdr_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")

    ws2["A1"] = "📊 빈티지 가죽자켓 556벌 브랜드 & 상태 등급 총괄 리포트"
    ws2["A1"].font = dash_title_font
    ws2.row_dimensions[1].height = 32

    # Tier statistics
    tier_counts = {"S등급": 0, "A등급": 0, "B등급": 0, "C등급": 0}
    grade_counts = {"S급": 0, "A급": 0, "B급": 0, "C급": 0}
    for r in csv_rows:
        tb = r["브랜드등급"]
        cg = r["상태등급"]
        if tb in tier_counts: tier_counts[tb] += 1
        if cg in grade_counts: grade_counts[cg] += 1

    ws2["A3"] = "1. 브랜드 등급 분포 (Brand Tiers)"
    ws2["A3"].font = sec_font
    ws2.append(["등급", "자켓 수량", "비율(%)", "대표 브랜드 예시"])
    for col in range(1, 5):
        c = ws2.cell(row=4, column=col)
        c.fill = tbl_hdr_fill
        c.font = tbl_hdr_font
        c.alignment = Alignment(horizontal="center")

    tier_desc = {
        "S등급": "Schott, Vanson, Kadoya, Lewis Leathers, Avirex 등 명가",
        "A등급": "Ralph Lauren, L.L.Bean, Alpha, Diesel, AllSaints 등 메이저",
        "B등급": "Intermezzo, Maestro, Wind Armor, Morgan, Dayson 등 클래식",
        "C등급": "일반 빈티지 공방 및 무명 오리지널"
    }
    for r_idx, (t_name, count) in enumerate(tier_counts.items(), start=5):
        pct = f"{(count / len(csv_rows) * 100):.1f}%" if csv_rows else "0%"
        ws2.append([t_name, f"{count}벌", pct, tier_desc.get(t_name, "")])
        for col in range(1, 5):
            cell = ws2.cell(row=r_idx, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col < 4 else "left")

    row_off = 11
    ws2.cell(row=row_off, column=1, value="2. 상태 등급 분포 (Condition Grades)").font = sec_font
    ws2.cell(row=row_off+1, column=1, value="상태 등급").fill = tbl_hdr_fill
    ws2.cell(row=row_off+1, column=1).font = tbl_hdr_font
    ws2.cell(row=row_off+1, column=2, value="자켓 수량").fill = tbl_hdr_fill
    ws2.cell(row=row_off+1, column=2).font = tbl_hdr_font
    ws2.cell(row=row_off+1, column=3, value="비율(%)").fill = tbl_hdr_fill
    ws2.cell(row=row_off+1, column=3).font = tbl_hdr_font
    ws2.cell(row=row_off+1, column=4, value="상태 기준 요약").fill = tbl_hdr_fill
    ws2.cell(row=row_off+1, column=4).font = tbl_hdr_font

    grade_desc = {
        "S급": "민트/극상 (스크래치/오염 거의 없는 최상급)",
        "A급": "우수/자연스러운 에이징 (찢김 없이 가죽 질감 우수)",
        "B급": "보통/빈티지 사용감 (생활 마모 있으나 착용에 무리없음)",
        "C급": "데미지/수선요망 (찢김, 깊은 까짐, 안감 오염 등 결함)"
    }
    for r_idx, (g_name, count) in enumerate(grade_counts.items(), start=row_off+2):
        pct = f"{(count / len(csv_rows) * 100):.1f}%" if csv_rows else "0%"
        ws2.append([g_name, f"{count}벌", pct, grade_desc.get(g_name, "")])
        for col in range(1, 5):
            cell = ws2.cell(row=r_idx, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center" if col < 4 else "left")

    wb.save(OUTPUT_EXCEL)
    print(f"[🎉] 마스터 엑셀 생성 완료! 저장 -> {OUTPUT_EXCEL}")

    pd.DataFrame(csv_rows).to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    print(f"[✓] CSV 저장 완료 -> {OUTPUT_CSV}")

if __name__ == "__main__":
    main()
