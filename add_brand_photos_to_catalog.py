import os, glob, sys, json, requests, uuid, re
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import pandas as pd
import base64, io

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets"
DOWNLOAD_DIR = os.path.join(BASE_DIR, "downloaded_jackets")
THUMB_DIR = os.path.join(BASE_DIR, "thumbnails")
os.makedirs(THUMB_DIR, exist_ok=True)

CACHE_CARDS_FILE = os.path.join(BASE_DIR, "detected_cards_cache_v2.json")
CHECKPOINT_FILE = os.path.join(BASE_DIR, "analyzed_jackets_checkpoint.json")
LABELS_CACHE_FILE = os.path.join(BASE_DIR, "jacket_labels_cache.json")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.xlsx")
OUTPUT_CSV = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.csv")

feynman_auth = os.path.expanduser('~/.feynman/agent/auth.json')
with open(feynman_auth, 'r', encoding='utf-8') as f:
    api_key = json.load(f)['opencode-go']['key']

def encode_thumb(img_path, max_dim=250):
    try:
        with Image.open(img_path) as im:
            im = im.convert('RGB')
            w, h = im.size
            if max(w, h) > max_dim:
                scale = max_dim / float(max(w, h))
                im = im.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
            buf = io.BytesIO()
            im.save(buf, format='JPEG', quality=75)
            return base64.b64encode(buf.getvalue()).decode('utf-8')
    except Exception:
        return None

def make_thumbnail(src_path, prefix, max_size=(75, 75)):
    if not src_path or not os.path.exists(src_path):
        return None
    fname = f"{prefix}_{os.path.basename(src_path)}"
    out_path = os.path.join(THUMB_DIR, fname)
    if os.path.exists(out_path):
        return out_path
    try:
        with Image.open(src_path) as im:
            im = im.convert("RGB")
            im.thumbnail(max_size, Image.Resampling.LANCZOS)
            im.save(out_path, "JPEG", quality=85)
        return out_path
    except Exception:
        return None

def find_label_worker(item):
    folder = item["folder"]
    folder_path = os.path.join(DOWNLOAD_DIR, folder)
    candidates = item["candidates"]

    if len(candidates) == 1:
        return item["key"], candidates[0]
    if len(candidates) == 0:
        return item["key"], item.get("front_photo", "")

    prompt = (
        "다음 사진들 중에서 자켓의 '브랜드 목 라벨'(브랜드 로고가 적힌 라벨 또는 목 안쪽 브랜드 상표/케어라벨) 사진 1개의 파일명만 JSON으로 출력하세요.\n"
        "예시: {\"brand_label_photo\": \"photo_006.jpg\"}"
    )

    content_list = [{"type": "text", "text": prompt}]
    for c in candidates:
        content_list.append({"type": "text", "text": c})
        p = os.path.join(folder_path, c)
        b64 = encode_thumb(p)
        if b64:
            content_list.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}})

    url = "https://opencode.ai/zen/go/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "x-opencode-session": f"ses_lbl_{uuid.uuid4().hex[:8]}"
    }
    payload = {
        "model": "deepseek-v4.1-flash",
        "messages": [{"role": "user", "content": content_list}],
        "max_tokens": 500
    }

    try:
        resp = requests.post(url, headers=headers, json=payload, timeout=30)
        if resp.status_code == 200:
            c = resp.json()["choices"][0]["message"]
            txt = c.get("content") or c.get("reasoning_content") or ""
            matches = re.findall(r"photo_\d+\.jpg", txt)
            valid = [m for m in matches if m in candidates]
            if valid:
                return item["key"], valid[0]
    except Exception as e:
        pass

    # Default heuristic: 2nd from last
    fallback = candidates[-2] if len(candidates) >= 2 else candidates[-1]
    return item["key"], fallback

def load_segments_with_photos():
    with open(CACHE_CARDS_FILE, "r", encoding="utf-8") as f:
        cache_data = json.load(f)

    with open(CHECKPOINT_FILE, "r", encoding="utf-8") as f:
        checkpoint = json.load(f)

    folders = sorted(glob.glob(os.path.join(DOWNLOAD_DIR, "*")))
    all_segs = []

    for fld in folders:
        fld_name = os.path.basename(fld)
        jpgs = sorted(glob.glob(os.path.join(fld, "photo_*.jpg")))
        if not jpgs or fld_name not in cache_data:
            continue

        item_title = fld_name
        info_file = os.path.join(fld, "item_info.txt")
        if os.path.exists(info_file):
            with open(info_file, "r", encoding="utf-8", errors="ignore") as inf:
                for l in inf:
                    if l.startswith("상품명:"): item_title = l.replace("상품명:", "").strip()

        cards = cache_data[fld_name]
        valid_cards = []
        for c in cards:
            c_path = os.path.join(fld, c)
            if os.path.exists(c_path) and os.path.getsize(c_path) != 33084:
                valid_cards.append(c)

        filtered_cards = []
        for c in valid_cards:
            num = int(c.replace("photo_", "").replace(".jpg", ""))
            if filtered_cards:
                prev_num = int(filtered_cards[-1].replace("photo_", "").replace(".jpg", ""))
                if num - prev_num < 3:
                    continue
            filtered_cards.append(c)

        fname_list = [os.path.basename(p) for p in jpgs]

        for idx, card_fname in enumerate(filtered_cards):
            card_idx = fname_list.index(card_fname) if card_fname in fname_list else -1
            if card_idx == -1: continue

            if idx + 1 < len(filtered_cards):
                next_card_idx = fname_list.index(filtered_cards[idx + 1]) if filtered_cards[idx + 1] in fname_list else len(fname_list)
            else:
                next_card_idx = len(fname_list)

            jkt_photos = fname_list[card_idx:next_card_idx]
            front_photo = jkt_photos[1] if len(jkt_photos) > 1 else card_fname

            # Candidate label photos: from index 3 onwards
            candidates = jkt_photos[3:] if len(jkt_photos) > 4 else jkt_photos[2:]

            key = f"{fld_name}_{card_fname}"
            seg_data = checkpoint.get(key, {})

            seg = {
                "key": key,
                "folder": fld_name,
                "jacket_index": idx + 1,
                "total_in_folder": len(filtered_cards),
                "card_photo": card_fname,
                "front_photo": front_photo,
                "candidates": candidates,
                "all_photos": jkt_photos,
                "photo_range": f"{jkt_photos[0]} ~ {jkt_photos[-1]} ({len(jkt_photos)}장)",
                "item_title": item_title,
                "jacket_code": seg_data.get("jacket_code", ""),
                "brand": seg_data.get("brand", "빈티지 오리지널"),
                "leather_type": seg_data.get("leather_type", "천연가죽"),
                "origin": seg_data.get("origin", "불명"),
                "shoulder_cm": seg_data.get("shoulder_cm", "-"),
                "chest_cm": seg_data.get("chest_cm", "-"),
                "length_cm": seg_data.get("length_cm", "-"),
                "sleeve_cm": seg_data.get("sleeve_cm", "-"),
            }
            all_segs.append(seg)

    return all_segs

def main():
    segments = load_segments_with_photos()
    print(f"[*] Total jackets: {len(segments)}")

    # Load label cache
    labels_cache = {}
    if os.path.exists(LABELS_CACHE_FILE):
        try:
            with open(LABELS_CACHE_FILE, "r", encoding="utf-8") as f:
                labels_cache = json.load(f)
            print(f"[*] Loaded {len(labels_cache)} labels from cache.")
        except Exception:
            labels_cache = {}

    to_process = [s for s in segments if s["key"] not in labels_cache]
    print(f"[*] Jackets remaining to detect label photos: {len(to_process)}")

    if to_process:
        with ThreadPoolExecutor(max_workers=12) as executor:
            future_to_item = {executor.submit(find_label_worker, item): item for item in to_process}
            count = 0
            for future in as_completed(future_to_item):
                count += 1
                try:
                    k, lbl = future.result()
                    labels_cache[k] = lbl
                    if count % 20 == 0 or count == len(to_process):
                        print(f"  [{count}/{len(to_process)}] Found label: {k} -> {lbl}", flush=True)
                        with open(LABELS_CACHE_FILE, "w", encoding="utf-8") as f:
                            json.dump(labels_cache, f, ensure_ascii=False, indent=2)
                except Exception as e:
                    print(f"Error: {e}")

        with open(LABELS_CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(labels_cache, f, ensure_ascii=False, indent=2)

    print("[+] All label photos identified! Now generating 3-photo Visual Excel...")

    # Build Excel with 3 photos: Card, Front, Brand Label
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "빈티지_가죽자켓_실물도감"

    headers = [
        "순번", "자켓품번", "실측카드 사진", "자켓 정면 사진", "브랜드 라벨 사진",
        "판독브랜드", "가죽소재", "원산지", 
        "어깨(cm)", "가슴(cm)", "기장(cm)", "소매(cm)", 
        "사진범위", "상품폴더", "상품명"
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 30

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = {
        1: 8,   # 순번
        2: 14,  # 자켓품번
        3: 15,  # 실측카드 사진 (C)
        4: 15,  # 자켓 정면 사진 (D)
        5: 15,  # 브랜드 라벨 사진 (E)
        6: 20,  # 판독브랜드
        7: 18,  # 가죽소재
        8: 15,  # 원산지
        9: 12,  # 어깨
        10: 12, # 가슴
        11: 12, # 기장
        12: 12, # 소매
        13: 24, # 사진범위
        14: 22, # 상품폴더
        15: 38  # 상품명
    }
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    csv_rows = []

    for idx, seg in enumerate(segments, start=1):
        label_photo = labels_cache.get(seg["key"], seg["candidates"][-1] if seg["candidates"] else "")
        code = seg.get("jacket_code") or f"{seg['folder'].replace('item_', '')}-{seg['jacket_index']}"
        brand = seg.get("brand") or "빈티지 오리지널"
        leather = seg.get("leather_type") or "천연가죽"
        origin = seg.get("origin") or "불명"
        sh = seg.get("shoulder_cm") or "-"
        ch = seg.get("chest_cm") or "-"
        ln = seg.get("length_cm") or "-"
        sl = seg.get("sleeve_cm") or "-"

        row_num = idx + 1
        ws.row_dimensions[row_num].height = 75

        row_vals = [
            idx, code, "", "", "", # C, D, E for embedded images
            brand, leather, origin,
            sh, ch, ln, sl,
            seg["photo_range"], seg["folder"], seg["item_title"]
        ]
        ws.append(row_vals)

        for col_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = Font(name="맑은 고딕", size=10)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 1. Column C: Embed Size Card Image
        thumb_card = make_thumbnail(os.path.join(DOWNLOAD_DIR, seg["folder"], seg["card_photo"]), f"{seg['folder']}_card")
        if thumb_card and os.path.exists(thumb_card):
            xl_c = XLImage(thumb_card)
            xl_c.width = 70
            xl_c.height = 70
            ws.add_image(xl_c, f"C{row_num}")

        # 2. Column D: Embed Jacket Front Image
        thumb_front = make_thumbnail(os.path.join(DOWNLOAD_DIR, seg["folder"], seg["front_photo"]), f"{seg['folder']}_front")
        if thumb_front and os.path.exists(thumb_front):
            xl_f = XLImage(thumb_front)
            xl_f.width = 70
            xl_f.height = 70
            ws.add_image(xl_f, f"D{row_num}")

        # 3. Column E: Embed Brand Label Image (NEW!)
        if label_photo:
            thumb_label = make_thumbnail(os.path.join(DOWNLOAD_DIR, seg["folder"], label_photo), f"{seg['folder']}_label")
            if thumb_label and os.path.exists(thumb_label):
                xl_l = XLImage(thumb_label)
                xl_l.width = 70
                xl_l.height = 70
                ws.add_image(xl_l, f"E{row_num}")

        csv_rows.append({
            "순번": idx,
            "자켓품번": code,
            "실측카드사진": seg["card_photo"],
            "자켓정면사진": seg["front_photo"],
            "브랜드라벨사진": label_photo,
            "판독브랜드": brand,
            "가죽소재": leather,
            "원산지": origin,
            "어깨(cm)": sh,
            "가슴(cm)": ch,
            "기장(cm)": ln,
            "소매(cm)": sl,
            "사진범위": seg["photo_range"],
            "상품폴더": seg["folder"],
            "상품명": seg["item_title"]
        })

    wb.save(OUTPUT_EXCEL)
    print(f"[🎉] Visual Excel with 3 photos created successfully at:\n  - {OUTPUT_EXCEL}")

    df_csv = pd.DataFrame(csv_rows)
    df_csv.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    print(f"[🎉] CSV created successfully at:\n  - {OUTPUT_CSV}")

if __name__ == "__main__":
    main()
