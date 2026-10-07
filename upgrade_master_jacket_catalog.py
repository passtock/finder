import os, glob, sys, json, requests, uuid, re
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import pandas as pd
import base64, io

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets"
DOWNLOAD_DIR = os.path.join(BASE_DIR, "downloaded_jackets")
THUMB_DIR = os.path.join(BASE_DIR, "thumbnails")
os.makedirs(THUMB_DIR, exist_ok=True)

CACHE_CARDS_FILE = os.path.join(BASE_DIR, "detected_cards_cache_v2.json")
CHECKPOINT_FILE = os.path.join(BASE_DIR, "analyzed_jackets_checkpoint.json")
LABELS_CACHE_FILE = os.path.join(BASE_DIR, "jacket_labels_cache.json")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.xlsx")
OUTPUT_CSV = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.csv")

feynman_auth = os.path.expanduser('~/.feynman/agent/auth.json')
with open(feynman_auth, 'r', encoding='utf-8') as f:
    api_key = json.load(f)['opencode-go']['key']

def encode_thumb(img_path, max_dim=250):
    try:
        with Image.open(img_path) as im:
            im = im.convert('RGB')
            w, h = im.size
            if max(w, h) > max_dim:
                scale = max_dim / float(max(w, h))
                im = im.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
            buf = io.BytesIO()
            im.save(buf, format='JPEG', quality=75)
            return base64.b64encode(buf.getvalue()).decode('utf-8')
    except Exception:
        return None

def make_thumbnail(src_path, prefix, max_size=(75, 75)):
    if not src_path or not os.path.exists(src_path):
        return None
    fname = f"{prefix}_{os.path.basename(src_path)}"
    out_path = os.path.join(THUMB_DIR, fname)
    if os.path.exists(out_path):
        return out_path
    try:
        with Image.open(src_path) as im:
            im = im.convert("RGB")
            im.thumbnail(max_size, Image.Resampling.LANCZOS)
            im.save(out_path, "JPEG", quality=85)
        return out_path
    except Exception:
        return None

def find_label_worker(item):
    folder = item["folder"]
    folder_path = os.path.join(DOWNLOAD_DIR, folder)
    candidates = item["candidates"]

    if len(candidates) == 1:
        return item["key"], candidates[0]
    if len(candidates) == 0:
        return item["key"], item.get("front_photo", "")

    prompt = (
        "다음 사진들 중에서 자켓의 '브랜드 목 라벨'(브랜드 로고가 적힌 라벨 또는 목 안쪽 브랜드 상표/케어라벨) 사진 1개의 파일명만 JSON으로 출력하세요.\n"
        "예시: {\"brand_label_photo\": \"photo_006.jpg\"}"
    )

    content_list = [{"type": "text", "text": prompt}]
    for c in candidates:
        content_list.append({"type": "text", "text": c})
        p = os.path.join(folder_path, c)
        b64 = encode_thumb(p)
        if b64:
            content_list.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}})

    url = "https://opencode.ai/zen/go/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "x-opencode-session": f"ses_lbl_{uuid.uuid4().hex[:8]}"
    }
    payload = {
        "model": "deepseek-v4.1-flash",
        "messages": [{"role": "user", "content": content_list}],
        "max_tokens": 500
    }

    try:
        resp = requests.post(url, headers=headers, json=payload, timeout=30)
        if resp.status_code == 200:
            c = resp.json()["choices"][0]["message"]
            txt = c.get("content") or c.get("reasoning_content") or ""
            matches = re.findall(r"photo_\d+\.jpg", txt)
            valid = [m for m in matches if m in candidates]
            if valid:
                return item["key"], valid[0]
    except Exception:
        pass

    fallback = candidates[-2] if len(candidates) >= 2 else candidates[-1]
    return item["key"], fallback

def analyze_jacket_deep_worker(seg, labels_cache):
    folder = seg["folder"]
    folder_path = os.path.join(DOWNLOAD_DIR, folder)
    card_p = os.path.join(folder_path, seg["card_photo"])
    front_p = os.path.join(folder_path, seg["front_photo"])
    label_fname = labels_cache.get(seg["key"], seg["front_photo"])
    label_p = os.path.join(folder_path, label_fname)

    prompt = (
        "당신은 전세계 빈티지 가죽자켓 및 일본 도메스틱/로컬 레더웨어 전문 감정가입니다.\n"
        "제공된 사진:\n"
        "1) [실측카드] (손글씨 사이즈표)\n"
        "2) [브랜드 라벨] (목 안쪽 상표 라벨, 케어라벨, 또는 가죽 패치)\n"
        "3) [자켓 정면 사진]\n\n"
        "[중요 판독 지침]:\n"
        "1. 브랜드명(brand): [브랜드 라벨] 사진에 적힌 영문/한자/일문 상표명을 돋보기로 보듯 정밀하게 읽어내세요.\n"
        "   - 유명 브랜드(Schott, Avirex 등)뿐만 아니라, 일본 도메스틱/로컬 브랜드, 가죽 공방(예: Wind Armor, Morgan Homme, Horn Works, Jackrose, Nicole Club, D'urban, Pazzo, Van Jacket, Intermezzo, Tete Homme, Abahouse, Person's, Liugoo, Freedom, G-Stage, Sears, Cooper 등)이라도 라벨에 적힌 글자 그대로 정확히 기재하세요.\n"
        "   - 라벨에 글자가 전혀 없고 순수 세탁기호나 Genuine Leather 마크만 있을 때만 '빈티지 오리지널 (무명)'으로 적으세요.\n"
        "2. 실측 치수: 실측카드에서 编号(jacket_code), 肩宽(shoulder_cm), 胸围(chest_cm), 衣长(length_cm), 袖长(sleeve_cm)을 정확한 숫자로 추출하세요.\n"
        "3. 가죽소재(leather_type) & 원산지(origin): 양가죽/소가죽/말가죽/돼지가죽/사슴가죽 등과 제조국(일본/미국/한국/이탈리아/파키스탄 등)을 추출하세요.\n\n"
        "반드시 순수 JSON으로만 응답하세요:\n"
        "{\n"
        '  "jacket_code": "자켓 품번 (예: M1, U12, O73 등)",\n'
        '  "brand": "판독된 브랜드명",\n'
        '  "leather_type": "가죽 종류",\n'
        '  "origin": "원산지",\n'
        '  "shoulder_cm": "어깨",\n'
        '  "chest_cm": "가슴",\n'
        '  "length_cm": "기장",\n'
        '  "sleeve_cm": "소매"\n'
        "}"
    )

    content_list = [{"type": "text", "text": prompt}]

    # 1) Card photo
    b64_card = encode_thumb(card_p, 300)
    if b64_card:
        content_list.append({"type": "text", "text": "[사진1: 손글씨 실측표 카드]"})
        content_list.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64_card}"}})

    # 2) Brand label photo
    b64_lbl = encode_thumb(label_p, 300)
    if b64_lbl:
        content_list.append({"type": "text", "text": "[사진2: 브랜드 목 라벨 / 상표]"})
        content_list.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64_lbl}"}})

    # 3) Front photo
    b64_front = encode_thumb(front_p, 250)
    if b64_front:
        content_list.append({"type": "text", "text": "[사진3: 자켓 정면 전체 사진]"})
        content_list.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64_front}"}})

    url = "https://opencode.ai/zen/go/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "x-opencode-session": f"ses_deep_{uuid.uuid4().hex[:8]}"
    }
    payload = {
        "model": "deepseek-v4.1-flash",
        "messages": [{"role": "user", "content": content_list}],
        "max_tokens": 2000
    }

    try:
        resp = requests.post(url, headers=headers, json=payload, timeout=35)
        if resp.status_code == 200:
            c = resp.json()["choices"][0]["message"]
            txt = c.get("content") or ""
            clean_txt = re.sub(r'```json\s*', '', txt)
            clean_txt = re.sub(r'```', '', clean_txt).strip()
            json_match = re.search(r"\{[\s\S]*\}", clean_txt)
            if json_match:
                raw_json = json_match.group(0)
                try:
                    parsed = json.loads(raw_json)
                except Exception:
                    cleaned_json = re.sub(r',\s*\}', '}', raw_json)
                    try:
                        parsed = json.loads(cleaned_json)
                    except Exception:
                        parsed = {}
                if parsed:
                    b = str(parsed.get("brand", "")).strip()
                    if not b or "라벨에 적힌" in b:
                        b = "빈티지 오리지널 (무명)"
                    parsed["brand"] = b
                    return seg["key"], parsed
    except Exception as e:
        pass

    return seg["key"], None

def load_all_segments():
    with open(CACHE_CARDS_FILE, "r", encoding="utf-8") as f:
        cache_data = json.load(f)

    folders = sorted(glob.glob(os.path.join(DOWNLOAD_DIR, "*")))
    all_segs = []

    for fld in folders:
        fld_name = os.path.basename(fld)
        jpgs = sorted(glob.glob(os.path.join(fld, "photo_*.jpg")))
        if not jpgs or fld_name not in cache_data:
            continue

        item_title = fld_name
        info_file = os.path.join(fld, "item_info.txt")
        if os.path.exists(info_file):
            with open(info_file, "r", encoding="utf-8", errors="ignore") as inf:
                for l in inf:
                    if l.startswith("상품명:"): item_title = l.replace("상품명:", "").strip()

        cards = cache_data[fld_name]
        valid_cards = []
        for c in cards:
            c_path = os.path.join(fld, c)
            if os.path.exists(c_path) and os.path.getsize(c_path) != 33084:
                valid_cards.append(c)

        filtered_cards = []
        for c in valid_cards:
            num = int(c.replace("photo_", "").replace(".jpg", ""))
            if filtered_cards:
                prev_num = int(filtered_cards[-1].replace("photo_", "").replace(".jpg", ""))
                if num - prev_num < 3:
                    continue
            filtered_cards.append(c)

        fname_list = [os.path.basename(p) for p in jpgs]

        for idx, card_fname in enumerate(filtered_cards):
            card_idx = fname_list.index(card_fname) if card_fname in fname_list else -1
            if card_idx == -1: continue

            if idx + 1 < len(filtered_cards):
                next_card_idx = fname_list.index(filtered_cards[idx + 1]) if filtered_cards[idx + 1] in fname_list else len(fname_list)
            else:
                next_card_idx = len(fname_list)

            jkt_photos = fname_list[card_idx:next_card_idx]
            front_photo = jkt_photos[1] if len(jkt_photos) > 1 else card_fname
            candidates = jkt_photos[3:] if len(jkt_photos) > 4 else jkt_photos[2:]

            key = f"{fld_name}_{card_fname}"

            seg = {
                "key": key,
                "folder": fld_name,
                "jacket_index": idx + 1,
                "total_in_folder": len(filtered_cards),
                "card_photo": card_fname,
                "front_photo": front_photo,
                "candidates": candidates,
                "all_photos": jkt_photos,
                "photo_range": f"{jkt_photos[0]} ~ {jkt_photos[-1]} ({len(jkt_photos)}장)",
                "item_title": item_title
            }
            all_segs.append(seg)

    return all_segs

def main():
    print("=" * 65)
    print("🧥 [Finder] 빈티지 가죽자켓 전체 도감 업그레이드 & 정밀 브랜드 판독기")
    print("=" * 65)

    segments = load_all_segments()
    print(f"[*] 총 탐지된 자켓 수: {len(segments)}벌")

    # 1. 브랜드 라벨 사진 캐시 로드 및 누락분 탐색
    labels_cache = {}
    if os.path.exists(LABELS_CACHE_FILE):
        try:
            with open(LABELS_CACHE_FILE, "r", encoding="utf-8") as f:
                labels_cache = json.load(f)
        except Exception:
            labels_cache = {}

    missing_labels = [s for s in segments if s["key"] not in labels_cache]
    print(f"[*] 라벨 사진 판별 필요: {len(missing_labels)}벌")

    if missing_labels:
        print("[*] 신규 자켓들의 브랜드 목 라벨 사진 찾는 중...")
        with ThreadPoolExecutor(max_workers=12) as executor:
            future_to_item = {executor.submit(find_label_worker, item): item for item in missing_labels}
            count = 0
            for future in as_completed(future_to_item):
                count += 1
                try:
                    k, lbl = future.result()
                    labels_cache[k] = lbl
                    if count % 20 == 0 or count == len(missing_labels):
                        print(f"  [라벨 탐색 {count}/{len(missing_labels)}] {k} -> {lbl}", flush=True)
                        with open(LABELS_CACHE_FILE, "w", encoding="utf-8") as f:
                            json.dump(labels_cache, f, ensure_ascii=False, indent=2)
                except Exception:
                    pass

        with open(LABELS_CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(labels_cache, f, ensure_ascii=False, indent=2)

    # 2. 체크포인트 로드
    checkpoint = {}
    if os.path.exists(CHECKPOINT_FILE):
        try:
            with open(CHECKPOINT_FILE, "r", encoding="utf-8") as f:
                checkpoint = json.load(f)
        except Exception:
            checkpoint = {}

    # 정밀 브랜드 재판독 대상 선별:
    # 1) 신규 자켓 (체크포인트에 없음)
    # 2) 기존 자켓 중 브랜드가 '빈티지 오리지널', '라벨에 적힌...', '' 인 경우 -> 실제 라벨 사진으로 정밀 재판독!
    to_deep_analyze = []
    for s in segments:
        k = s["key"]
        if k not in checkpoint:
            to_deep_analyze.append(s)
        else:
            prev_b = str(checkpoint[k].get("brand", "")).strip()
            if prev_b in ["빈티지 오리지널", "빈티지 오리지널 (미상)", ""] or "라벨에 적힌" in prev_b:
                to_deep_analyze.append(s)

    print(f"\n[*] 정밀 브랜드 & 치수 DeepSeek 판독 대상: {len(to_deep_analyze)}벌")
    print(f"    (신규 자켓 및 기존 '빈티지 오리지널' 미상 항목 전체 재검토 포함)")

    if to_deep_analyze:
        with ThreadPoolExecutor(max_workers=16) as executor:
            future_to_item = {executor.submit(analyze_jacket_deep_worker, s, labels_cache): s for s in to_deep_analyze}
            count = 0
            upgraded_brands = 0
            for future in as_completed(future_to_item):
                count += 1
                try:
                    k, data = future.result()
                    if data:
                        if k in checkpoint:
                            merged = checkpoint[k].copy()
                            for field in ["brand", "leather_type", "origin", "shoulder_cm", "chest_cm", "length_cm", "sleeve_cm", "jacket_code"]:
                                val = str(data.get(field, "")).strip()
                                if val and val not in ["-", "불명", "미상", "빈티지 오리지널 (무명)"]:
                                    merged[field] = val
                                elif field == "brand" and val:
                                    merged["brand"] = val
                            checkpoint[k] = merged
                        else:
                            checkpoint[k] = data
                        brand_name = checkpoint[k].get("brand", "")
                        if brand_name and "빈티지 오리지널" not in brand_name:
                            upgraded_brands += 1
                    if count % 15 == 0 or count == len(to_deep_analyze):
                        print(f"  [AI 정밀 판독 {count}/{len(to_deep_analyze)}] {k} -> 브랜드: {checkpoint.get(k, {}).get('brand', '불명')}, 품번: {checkpoint.get(k, {}).get('jacket_code', '-')}", flush=True)
                        with open(CHECKPOINT_FILE, "w", encoding="utf-8") as f:
                            json.dump(checkpoint, f, ensure_ascii=False, indent=2)
                except Exception as e:
                    pass

        with open(CHECKPOINT_FILE, "w", encoding="utf-8") as f:
            json.dump(checkpoint, f, ensure_ascii=False, indent=2)
        print(f"[✓] AI 판독 완료! 새롭게 찾아낸 브랜드: {upgraded_brands}개")

    # 3. 3-Photo 마스터 엑셀 및 CSV 빌드
    print("\n[*] [실측카드, 정면, 브랜드라벨] 3종 실물 사진 임베딩 마스터 엑셀 생성 중...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "빈티지_가죽자켓_총망라도감"

    headers = [
        "순번", "자켓품번", "실측카드 사진", "자켓 정면 사진", "브랜드 라벨 사진",
        "판독브랜드", "가죽소재", "원산지", 
        "어깨(cm)", "가슴(cm)", "기장(cm)", "소매(cm)", 
        "사진범위", "상품폴더", "상품명"
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="맑은 고딕", size=10)
    bold_font = Font(name="맑은 고딕", size=10, bold=True)
    brand_font = Font(name="맑은 고딕", size=10, bold=True, color="1E3A8A")

    csv_rows = []

    for idx, seg in enumerate(segments, start=1):
        row_num = idx + 1
        ws.row_dimensions[row_num].height = 65

        k = seg["key"]
        seg_data = checkpoint.get(k, {})
        folder_path = os.path.join(DOWNLOAD_DIR, seg["folder"])

        card_path = os.path.join(folder_path, seg["card_photo"])
        front_path = os.path.join(folder_path, seg["front_photo"])
        lbl_fname = labels_cache.get(k, seg["front_photo"])
        lbl_path = os.path.join(folder_path, lbl_fname)

        code = seg_data.get("jacket_code", "")
        if not code or code == "-":
            # 폴더명이나 상품명에서 힌트 추출
            m_c = re.search(r'[A-Za-z]\d{1,4}', seg["item_title"])
            code = f"{m_c.group(0)}-#{seg['jacket_index']}" if m_c else f"J{idx:03d}"

        brand = seg_data.get("brand", "빈티지 오리지널")
        leather = seg_data.get("leather_type", "천연가죽")
        origin = seg_data.get("origin", "불명")
        shoulder = seg_data.get("shoulder_cm", "-")
        chest = seg_data.get("chest_cm", "-")
        length = seg_data.get("length_cm", "-")
        sleeve = seg_data.get("sleeve_cm", "-")

        row_data = [
            idx,
            code,
            "", # 실측카드
            "", # 정면
            "", # 브랜드라벨
            brand,
            leather,
            origin,
            shoulder,
            chest,
            length,
            sleeve,
            seg["photo_range"],
            seg["folder"],
            seg["item_title"]
        ]
        ws.append(row_data)

        fill_color = zebra_fill if idx % 2 == 0 else white_fill
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=c_idx)
            cell.fill = fill_color
            cell.border = thin_border
            cell.font = data_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=row_num, column=2).font = bold_font
        ws.cell(row=row_num, column=6).font = brand_font
        ws.cell(row=row_num, column=15).alignment = Alignment(horizontal="left", vertical="center")

        # 3종 실물 사진 임베딩 (Col C, D, E)
        # 1. 실측카드 (Col C)
        t_card = make_thumbnail(card_path, f"card_{seg['folder']}")
        if t_card and os.path.exists(t_card):
            try:
                img_c = XLImage(t_card)
                ws.add_image(img_c, f"C{row_num}")
            except Exception: pass

        # 2. 정면 사진 (Col D)
        t_front = make_thumbnail(front_path, f"front_{seg['folder']}")
        if t_front and os.path.exists(t_front):
            try:
                img_f = XLImage(t_front)
                ws.add_image(img_f, f"D{row_num}")
            except Exception: pass

        # 3. 브랜드 라벨 사진 (Col E)
        t_lbl = make_thumbnail(lbl_path, f"lbl_{seg['folder']}")
        if t_lbl and os.path.exists(t_lbl):
            try:
                img_l = XLImage(t_lbl)
                ws.add_image(img_l, f"E{row_num}")
            except Exception: pass

        # CSV row
        csv_rows.append({
            "순번": idx,
            "자켓품번": code,
            "실측카드사진": os.path.basename(card_path),
            "자켓정면사진": os.path.basename(front_path),
            "브랜드라벨사진": os.path.basename(lbl_path),
            "판독브랜드": brand,
            "가죽소재": leather,
            "원산지": origin,
            "어깨(cm)": shoulder,
            "가슴(cm)": chest,
            "기장(cm)": length,
            "소매(cm)": sleeve,
            "사진범위": seg["photo_range"],
            "상품폴더": seg["folder"],
            "상품명": seg["item_title"]
        })

    # Column widths
    widths = {
        "A": 7, "B": 13, "C": 14, "D": 14, "E": 14,
        "F": 20, "G": 14, "H": 12, "I": 10, "J": 10,
        "K": 10, "L": 10, "M": 26, "N": 20, "O": 38
    }
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w

    ws.freeze_panes = "F2"
    wb.save(OUTPUT_EXCEL)
    print(f"[🎉] 총 {len(segments)}벌 3종 사진 임베딩 완료! 엑셀 저장 -> {OUTPUT_EXCEL}")

    pd.DataFrame(csv_rows).to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    print(f"[✓] CSV 저장 완료 -> {OUTPUT_CSV}")

if __name__ == "__main__":
    main()
