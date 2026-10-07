import os, sys
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets\test_visual_excel.xlsx"
thumb_dir = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets\thumbnails"
os.makedirs(thumb_dir, exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "가죽자켓_실물사진_목록"

headers = [
    "순번", "자켓품번", "실측카드 사진", "자켓 정면 사진", 
    "판독브랜드", "가죽소재", "원산지", 
    "어깨(cm)", "가슴(cm)", "기장(cm)", "소매(cm)", 
    "사진범위", "상품폴더"
]

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws.append(headers)
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

ws.row_dimensions[1].height = 28

# Sample 3 test jackets from item_1081695145771
folder_base = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets\downloaded_jackets\item_1081695145771"
sample_data = [
    {
        "seq": 1, "code": "P211", "card_img": os.path.join(folder_base, "photo_002.jpg"),
        "jacket_img": os.path.join(folder_base, "photo_003.jpg"),
        "brand": "빈티지 오리지널", "leather": "양가죽", "origin": "한국",
        "shoulder": "45", "chest": "50", "length": "72", "sleeve": "61",
        "range": "photo_002 ~ photo_008", "folder": "item_1081695145771"
    },
    {
        "seq": 2, "code": "P212", "card_img": os.path.join(folder_base, "photo_009.jpg"),
        "jacket_img": os.path.join(folder_base, "photo_010.jpg"),
        "brand": "빈티지 오리지널", "leather": "양가죽", "origin": "한국",
        "shoulder": "49", "chest": "53", "length": "70", "sleeve": "62",
        "range": "photo_009 ~ photo_015", "folder": "item_1081695145771"
    },
    {
        "seq": 3, "code": "P213", "card_img": os.path.join(folder_base, "photo_016.jpg"),
        "jacket_img": os.path.join(folder_base, "photo_017.jpg"),
        "brand": "빈티지 오리지널", "leather": "양가죽", "origin": "한국",
        "shoulder": "49", "chest": "53", "length": "76", "sleeve": "61",
        "range": "photo_016 ~ photo_022", "folder": "item_1081695145771"
    }
]

def make_thumb(src_path, prefix, max_size=(85, 85)):
    out_path = os.path.join(thumb_dir, f"{prefix}_{os.path.basename(src_path)}")
    with Image.open(src_path) as im:
        im = im.convert("RGB")
        im.thumbnail(max_size, Image.Resampling.LANCZOS)
        im.save(out_path, "JPEG", quality=85)
    return out_path

for r_idx, d in enumerate(sample_data, start=2):
    row_vals = [
        d["seq"], d["code"], "", "", # C and D are for images
        d["brand"], d["leather"], d["origin"],
        d["shoulder"], d["chest"], d["length"], d["sleeve"],
        d["range"], d["folder"]
    ]
    ws.append(row_vals)
    ws.row_dimensions[r_idx].height = 75 # Set row height for thumbnails
    
    for c_idx in range(1, len(row_vals) + 1):
        cell = ws.cell(row=r_idx, column=c_idx)
        cell.font = Font(name="맑은 고딕", size=10)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 1. Insert size card thumbnail into Col C (Col 3)
    if os.path.exists(d["card_img"]):
        t_card = make_thumb(d["card_img"], "card")
        xl_c = XLImage(t_card)
        xl_c.width = 70
        xl_c.height = 70
        ws.add_image(xl_c, f"C{r_idx}")
        
    # 2. Insert jacket front thumbnail into Col D (Col 4)
    if os.path.exists(d["jacket_img"]):
        t_jkt = make_thumb(d["jacket_img"], "jkt")
        xl_j = XLImage(t_jkt)
        xl_j.width = 70
        xl_j.height = 70
        ws.add_image(xl_j, f"D{r_idx}")

# Widths
widths = {'A': 8, 'B': 12, 'C': 14, 'D': 14, 'E': 16, 'F': 12, 'G': 10, 'H': 11, 'I': 11, 'J': 11, 'K': 11, 'L': 24, 'M': 22}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

wb.save(excel_path)
print(f"Sample visual Excel successfully created at: {excel_path}")
