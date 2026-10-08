import os, glob, sys, json, re, time
from PIL import Image
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets"
DOWNLOAD_DIR = os.path.join(BASE_DIR, "downloaded_jackets")
THUMB_DIR = os.path.join(BASE_DIR, "thumbnails")
os.makedirs(THUMB_DIR, exist_ok=True)

RESEARCH_FILE = os.path.join(BASE_DIR, "brand_research_database.json")
EVAL_CHECKPOINT_FILE = os.path.join(BASE_DIR, "jacket_evaluations_checkpoint.json")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.xlsx")
OUTPUT_CSV = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.csv")

with open(RESEARCH_FILE, 'r', encoding='utf-8') as f:
    brand_research = json.load(f)

with open(EVAL_CHECKPOINT_FILE, 'r', encoding='utf-8') as f:
    eval_checkpoint = json.load(f)

from upgrade_master_jacket_catalog import load_all_segments, make_thumbnail, LABELS_CACHE_FILE

segments = load_all_segments()
labels_cache = {}
if os.path.exists(LABELS_CACHE_FILE):
    with open(LABELS_CACHE_FILE, "r", encoding="utf-8") as f:
        labels_cache = json.load(f)

print(f"[*] 총 자켓 수: {len(segments)}벌")
print(f"[*] 브랜드 연구 데이터: {len(brand_research)}개 브랜드 등록됨")

# Function to match a jacket's brand to researched database
def match_research(jacket_brand_str):
    if not jacket_brand_str:
        return brand_research.get("빈티지 오리지널 (무명)")
    
    j_clean = jacket_brand_str.strip()
    unnamed_keywords = ['무명', '판독 불가', '판독불가', '미상', '미식별', '미표기', 
                        '식별 불가', '식별불가', '라벨 미확인', '미확인', '불명', 
                        '무지라벨', '블랭크', '빈티지 오리지널', 'Non-brand', 'non-brand', '식별불능']
    if any(w in j_clean for w in unnamed_keywords):
        return brand_research.get("빈티지 오리지널 (무명)")
    
    # 1. Exact match in research keys
    if j_clean in brand_research:
        return brand_research[j_clean]
        
    # 2. Case-insensitive substring match
    for b_key, b_info in brand_research.items():
        if b_key != "빈티지 오리지널 (무명)" and (b_key.lower() in j_clean.lower() or j_clean.lower() in b_key.lower()):
            return b_info
            
    # 3. Fallback
    return {
        "official_name": j_clean,
        "country": "불명",
        "category": "로컬 빈티지 레더웨어",
        "heritage": "각국의 로컬 피혁 공방 및 도메스틱 패션 브랜드에서 생산된 천연가죽 의류.",
        "leather_specialty": "천연 양가죽 / 소가죽",
        "retail_tier": "신품 25만~50만원대 / 빈티지 8만~18만원대",
        "tier": "C등급 (일반 빈티지 공방/도메스틱)"
    }

# Process and enrich each segment
enriched_segments = []
for seg in segments:
    k = seg["key"]
    ev = eval_checkpoint.get(k, {})
    
    raw_b = ev.get("brand_name", "")
    res = match_research(raw_b)
    
    b_official = res["official_name"]
    b_country = res["country"]
    b_tier_full = res["tier"]
    b_tier_code = b_tier_full.split()[0]  # S등급, A등급, B등급, C등급
    b_heritage = res["heritage"]
    b_retail = res.get("retail_tier", "")
    
    c_grade = ev.get("condition_grade", "B급")
    if not c_grade.endswith("급"): c_grade = f"{c_grade[0]}급"
    c_details = ev.get("condition_details", "자연스러운 빈티지 에이징 상태")
    
    # Calculate research-backed recommendation
    if b_tier_code == "S등급":
        if c_grade in ["S급", "A급"]: rec = "⭐⭐⭐⭐⭐ 강력추천(S)"
        elif c_grade == "B급": rec = "⭐⭐⭐⭐ 추천(A)"
        else: rec = "⭐⭐ 비추천(C)"
    elif b_tier_code == "A등급":
        if c_grade in ["S급", "A급"]: rec = "⭐⭐⭐⭐ 추천(A)"
        elif c_grade == "B급": rec = "⭐⭐⭐ 보통(B)"
        else: rec = "⭐⭐ 비추천(C)"
    elif b_tier_code == "B등급":
        if c_grade in ["S급", "A급"]: rec = "⭐⭐⭐⭐ 추천(A)"
        elif c_grade == "B급": rec = "⭐⭐⭐ 보통(B)"
        else: rec = "⭐⭐ 비추천(C)"
    else: # C등급 (무명 / 공방)
        if c_grade in ["S급", "A급"]: rec = "⭐⭐⭐ 보통(B)"
        elif c_grade == "B급": rec = "⭐⭐⭐ 보통(B)"
        else: rec = "⭐⭐ 비추천(C)"
        
    # Update evaluation checkpoint object
    ev["brand_name"] = b_official
    ev["brand_country"] = b_country
    ev["brand_tier"] = b_tier_code
    ev["brand_heritage"] = b_heritage
    ev["brand_retail"] = b_retail
    ev["condition_grade"] = c_grade
    ev["overall_recommendation"] = rec
    eval_checkpoint[k] = ev
    
    enriched_segments.append({
        "seg": seg,
        "key": k,
        "code": ev.get("jacket_code", f"J#{seg['jacket_index']}"),
        "brand": b_official,
        "country": b_country,
        "tier": b_tier_code,
        "heritage": f"{b_heritage} [출시가: {b_retail}]" if b_retail else b_heritage,
        "condition_grade": c_grade,
        "condition_details": c_details,
        "recommendation": rec,
        "leather_type": ev.get("leather_type", "천연가죽"),
        "origin": ev.get("origin", b_country),
        "shoulder": ev.get("shoulder_cm", "-"),
        "chest": ev.get("chest_cm", "-"),
        "length": ev.get("length_cm", "-"),
        "sleeve": ev.get("sleeve_cm", "-"),
        "selected_photos": ev.get("selected_photos", {})
    })

# Save updated checkpoint
with open(EVAL_CHECKPOINT_FILE, "w", encoding="utf-8") as f:
    json.dump(eval_checkpoint, f, ensure_ascii=False, indent=2)
print(f"[✓] {EVAL_CHECKPOINT_FILE} 체크포인트 연구 데이터 업데이트 완료")

# Sort priority: S tier & Recommended first
def sort_priority(item):
    score = 0
    if "강력추천(S)" in item["recommendation"]: score += 100
    elif "추천(A)" in item["recommendation"]: score += 70
    elif "보통(B)" in item["recommendation"]: score += 40
    else: score += 10
    
    if item["tier"] == "S등급": score += 20
    elif item["tier"] == "A등급": score += 15
    elif item["tier"] == "B등급": score += 10
    
    if item["condition_grade"] in ["S급", "A급"]: score += 5
    return -score

sorted_items = sorted(enriched_segments, key=sort_priority)

# Build Master Excel
print("\n[*] 마스터 엑셀 도감 빌드 중 (연구 기반 브랜드 프로필 & 5종 실물 사진 포함)...")
wb = openpyxl.Workbook()

# Sheet 1: Master Catalog
ws1 = wb.active
ws1.title = "자켓_종합평가_도감"

headers1 = [
    "순번", "자켓품번", "종합추천", "브랜드 등급", "상태 등급", 
    "판독 브랜드", "브랜드 국적", "브랜드 역사 & 리테일 조사 내용", "상태 상세평가 코멘트",
    "실측카드 사진", "자켓 정면 사진", "자켓 뒷면 사진", "브랜드 라벨 사진", "디테일/안감 사진",
    "가죽 소재", "원산지", "어깨(cm)", "가슴(cm)", "기장(cm)", "소매(cm)",
    "전체 사진수", "사진 범위", "상품폴더", "상품명"
]

header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)

ws1.append(headers1)
for col_idx in range(1, len(headers1) + 1):
    c = ws1.cell(row=1, column=col_idx)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws1.row_dimensions[1].height = 30

tier_colors = {
    "S등급": {"bg": "FEF3C7", "fg": "92400E", "bold": True},  # Gold
    "A등급": {"bg": "DCFCE7", "fg": "166534", "bold": True},  # Emerald
    "B등급": {"bg": "E0F2FE", "fg": "075985", "bold": True},  # Sky
    "C등급": {"bg": "F3F4F6", "fg": "4B5563", "bold": False}   # Gray
}
grade_colors = {
    "S급": {"bg": "FEF3C7", "fg": "92400E", "bold": True},
    "A급": {"bg": "DCFCE7", "fg": "166534", "bold": True},
    "B급": {"bg": "E0F2FE", "fg": "075985", "bold": True},
    "C급": {"bg": "FEE2E2", "fg": "991B1B", "bold": True}
}

zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
data_font = Font(name="맑은 고딕", size=9)
bold_font = Font(name="맑은 고딕", size=9, bold=True)
desc_font = Font(name="맑은 고딕", size=8)

csv_rows = []

for idx, itm in enumerate(sorted_items, start=1):
    row_num = idx + 1
    ws1.row_dimensions[row_num].height = 68

    seg = itm["seg"]
    fld_path = os.path.join(DOWNLOAD_DIR, seg["folder"])

    sel_photos = itm["selected_photos"]
    card_p = os.path.join(fld_path, sel_photos.get("card", seg["card_photo"]))
    front_p = os.path.join(fld_path, sel_photos.get("front", seg["front_photo"]))
    back_p = os.path.join(fld_path, sel_photos.get("back", seg["all_photos"][2] if len(seg["all_photos"])>2 else seg["front_photo"]))
    label_p = os.path.join(fld_path, sel_photos.get("label", labels_cache.get(seg["key"], seg["front_photo"])))
    detail_p = os.path.join(fld_path, sel_photos.get("detail", seg["all_photos"][-1] if len(seg["all_photos"])>4 else seg["front_photo"]))

    row_data = [
        idx,
        itm["code"],
        itm["recommendation"],
        itm["tier"],
        itm["condition_grade"],
        itm["brand"],
        itm["country"],
        itm["heritage"],
        itm["condition_details"],
        "", "", "", "", "", # 5 embedded photos
        itm["leather_type"],
        itm["origin"],
        itm["shoulder"],
        itm["chest"],
        itm["length"],
        itm["sleeve"],
        f"{len(seg['all_photos'])}장",
        seg["photo_range"],
        seg["folder"],
        seg["item_title"]
    ]
    ws1.append(row_data)

    fill_color = zebra_fill if idx % 2 == 0 else white_fill
    for c_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=row_num, column=c_idx)
        cell.fill = fill_color
        cell.border = thin_border
        cell.font = data_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws1.cell(row=row_num, column=2).font = bold_font
    ws1.cell(row=row_num, column=6).font = Font(name="맑은 고딕", size=9, bold=True, color="1E3A8A")
    ws1.cell(row=row_num, column=8).font = desc_font
    ws1.cell(row=row_num, column=8).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws1.cell(row=row_num, column=9).font = desc_font
    ws1.cell(row=row_num, column=9).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws1.cell(row=row_num, column=24).alignment = Alignment(horizontal="left", vertical="center")

    # Tier badge styling
    if itm["tier"] in tier_colors:
        tc = tier_colors[itm["tier"]]
        ws1.cell(row=row_num, column=4).fill = PatternFill(start_color=tc["bg"], end_color=tc["bg"], fill_type="solid")
        ws1.cell(row=row_num, column=4).font = Font(name="맑은 고딕", size=9, bold=tc["bold"], color=tc["fg"])

    # Condition badge styling
    if itm["condition_grade"] in grade_colors:
        gc = grade_colors[itm["condition_grade"]]
        ws1.cell(row=row_num, column=5).fill = PatternFill(start_color=gc["bg"], end_color=gc["bg"], fill_type="solid")
        ws1.cell(row=row_num, column=5).font = Font(name="맑은 고딕", size=9, bold=gc["bold"], color=gc["fg"])

    # Embed 5 Real Thumbnails (Cols J, K, L, M, N)
    photo_targets = [
        (card_p, "J", f"card_{seg['folder']}"),
        (front_p, "K", f"front_{seg['folder']}"),
        (back_p, "L", f"back_{seg['folder']}"),
        (label_p, "M", f"lbl_{seg['folder']}"),
        (detail_p, "N", f"det_{seg['folder']}")
    ]

    for p_src, col_letter, p_prefix in photo_targets:
        t_path = make_thumbnail(p_src, p_prefix)
        if t_path and os.path.exists(t_path):
            try:
                img = XLImage(t_path)
                ws1.add_image(img, f"{col_letter}{row_num}")
            except Exception:
                pass

    csv_rows.append({
        "순번": idx,
        "자켓품번": itm["code"],
        "종합추천": itm["recommendation"],
        "브랜드등급": itm["tier"],
        "상태등급": itm["condition_grade"],
        "판독브랜드": itm["brand"],
        "브랜드국적": itm["country"],
        "브랜드조사내용": itm["heritage"],
        "상태상세평가": itm["condition_details"],
        "가죽소재": itm["leather_type"],
        "원산지": itm["origin"],
        "어깨(cm)": itm["shoulder"],
        "가슴(cm)": itm["chest"],
        "기장(cm)": itm["length"],
        "소매(cm)": itm["sleeve"],
        "실측카드사진": os.path.basename(card_p),
        "자켓정면사진": os.path.basename(front_p),
        "자켓뒷면사진": os.path.basename(back_p),
        "브랜드라벨사진": os.path.basename(label_p),
        "디테일안감사진": os.path.basename(detail_p),
        "전체사진수": len(seg["all_photos"]),
        "사진범위": seg["photo_range"],
        "상품폴더": seg["folder"],
        "상품명": seg["item_title"]
    })

# Widths for Sheet 1
widths1 = {
    "A": 6, "B": 11, "C": 18, "D": 11, "E": 11,
    "F": 20, "G": 14, "H": 46, "I": 40,
    "J": 13, "K": 13, "L": 13, "M": 13, "N": 13,
    "O": 14, "P": 10, "Q": 9, "R": 9, "S": 9, "T": 9,
    "U": 10, "V": 26, "W": 20, "X": 35
}
for col_l, w in widths1.items():
    ws1.column_dimensions[col_l].width = w

ws1.freeze_panes = "J2"

# Sheet 2: Researched Brand Directory & Dashboard
ws2 = wb.create_sheet(title="브랜드별_정밀조사_디렉토리")
ws2.column_dimensions["A"].width = 24
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 12
ws2.column_dimensions["D"].width = 10
ws2.column_dimensions["E"].width = 28
ws2.column_dimensions["F"].width = 50
ws2.column_dimensions["G"].width = 30

dash_title_font = Font(name="맑은 고딕", size=14, bold=True, color="1E3A8A")
sec_font = Font(name="맑은 고딕", size=11, bold=True, color="1F2937")
tbl_hdr_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
tbl_hdr_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")

ws2["A1"] = "📚 전수 추출 브랜드 정밀 조사 디렉토리 & 가치 평가표"
ws2["A1"].font = dash_title_font
ws2.row_dimensions[1].height = 32

headers2 = ["브랜드명", "국적", "브랜드 등급", "보유수량", "카테고리 및 분류", "브랜드 역사 & 가죽 특성", "출시 및 빈티지 시세"]
ws2.append(headers2)
for c_idx in range(1, len(headers2) + 1):
    c = ws2.cell(row=2, column=c_idx)
    c.fill = tbl_hdr_fill
    c.font = tbl_hdr_font
    c.alignment = Alignment(horizontal="center", vertical="center")
ws2.row_dimensions[2].height = 26

# Sort brands by Tier (S -> A -> B -> C) then count
def sort_brand_key(item):
    b_name, info = item
    t = info.get("tier", "C등급").split()[0]
    rank = {"S등급": 4, "A등급": 3, "B등급": 2, "C등급": 1}.get(t, 0)
    cnt = info.get("jacket_count", 0)
    return (-rank, -cnt)

sorted_brand_items = sorted(brand_research.items(), key=sort_brand_key)

for r_idx, (b_name, info) in enumerate(sorted_brand_items, start=3):
    ws2.row_dimensions[r_idx].height = 30
    t_code = info.get("tier", "C등급").split()[0]
    cnt = info.get("jacket_count", 0)
    row_vals = [
        info.get("official_name", b_name),
        info.get("country", "-"),
        t_code,
        f"{cnt}벌",
        info.get("category", "-"),
        info.get("heritage", "-"),
        info.get("retail_tier", "-")
    ]
    ws2.append(row_vals)
    
    fill = zebra_fill if r_idx % 2 == 0 else white_fill
    for c_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=r_idx, column=c_idx)
        cell.fill = fill
        cell.border = thin_border
        cell.font = data_font
        cell.alignment = Alignment(horizontal="center" if c_idx in [2, 3, 4] else "left", vertical="center")
        
    if t_code in tier_colors:
        tc = tier_colors[t_code]
        ws2.cell(row=r_idx, column=3).fill = PatternFill(start_color=tc["bg"], end_color=tc["bg"], fill_type="solid")
        ws2.cell(row=r_idx, column=3).font = Font(name="맑은 고딕", size=9, bold=tc["bold"], color=tc["fg"])

ws2.freeze_panes = "A3"

# Save workbook and CSV
wb.save(OUTPUT_EXCEL)
print(f"[🎉] 연구 데이터가 반영된 마스터 엑셀 생성 완료! 저장 -> {OUTPUT_EXCEL}")

pd.DataFrame(csv_rows).to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
print(f"[✓] CSV 저장 완료 -> {OUTPUT_CSV}")
