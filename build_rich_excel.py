import os, sys
import pandas as pd
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')

base = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets\downloaded_jackets"
csv_path = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets\taobao_leather_jackets_with_brands.csv"
xlsx_path = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets\taobao_leather_jackets_with_brands.xlsx"

df = pd.read_csv(csv_path, encoding='utf-8-sig')

# Fix typos discovered during vision check
df.loc[df['자켓품번'] == 'Q911', '자켓품번'] = 'Q91'
df.loc[df['자켓품번'] == 'V180', '자켓품번'] = 'V100'

# Add local full path and hyperlink column
df['실측사진경로'] = df.apply(lambda r: os.path.join(base, r['상품폴더'], r['실측사진']), axis=1)

df.to_csv(csv_path, index=False, encoding='utf-8-sig')
print("Updated CSV saved.")

# Create rich Excel with hyperlinks and embedded image thumbnails
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "타오바오_가죽자켓_브랜드실측"

headers = [
    "순번", "자켓품번", "판독브랜드", "가죽소재", "원산지", 
    "어깨(cm)", "가슴(cm)", "기장(cm)", "소매(cm)", 
    "실측카드 미리보기", "실측사진 파일열기", "총사진수", "상품폴더", "상품명"
]

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws.append(headers)
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Thumbnails scratch directory
thumb_dir = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets\thumbnails"
os.makedirs(thumb_dir, exist_ok=True)

for row_idx, r in df.iterrows():
    excel_row = row_idx + 2
    img_path = r['실측사진경로']
    
    row_data = [
        r['순번'], r['자켓품번'], r['판독브랜드'], r['가죽소재'], r['원산지'],
        r['어깨(cm)'], r['가슴(cm)'], r['기장(cm)'], r['소매(cm)'],
        "", # Image cell placeholder (Col J)
        f'=HYPERLINK("{img_path}", "{r["실측사진"]}")', # Clickable link (Col K)
        r['총사진수'], r['상품폴더'], r['상품명']
    ]
    ws.append(row_data)
    ws.row_dimensions[excel_row].height = 80 # Expand row height for thumbnail
    
    # Format cells
    for col_num in range(1, len(row_data) + 1):
        c = ws.cell(row=excel_row, column=col_num)
        c.font = Font(name="맑은 고딕", size=10)
        c.border = thin_border
        if col_num in [1, 2, 5, 6, 7, 8, 9, 11, 12]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        elif col_num == 10:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")

    # Insert Thumbnail into Col J (Col 10)
    if os.path.exists(img_path):
        try:
            thumb_path = os.path.join(thumb_dir, f"thumb_{r['상품폴더']}.jpg")
            with Image.open(img_path) as im:
                im = im.convert("RGB")
                im.thumbnail((100, 100), Image.Resampling.LANCZOS)
                im.save(thumb_path, "JPEG", quality=85)
            
            xl_img = XLImage(thumb_path)
            xl_img.width = 95
            xl_img.height = 95
            cell_coord = f"J{excel_row}"
            ws.add_image(xl_img, cell_coord)
        except Exception as e:
            print(f"Error adding thumbnail for row {excel_row}: {e}")

# Column widths
col_widths = {
    'A': 8, 'B': 12, 'C': 16, 'D': 14, 'E': 10,
    'F': 12, 'G': 12, 'H': 12, 'I': 12,
    'J': 16, 'K': 18, 'L': 12, 'M': 22, 'N': 50
}
for col, width in col_widths.items():
    ws.column_dimensions[col].width = width

wb.save(xlsx_path)
print(f"Upgraded Excel saved with embedded thumbnails & links: {xlsx_path}")
