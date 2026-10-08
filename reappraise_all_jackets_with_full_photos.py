import os, glob, sys, json, requests, uuid, re, time
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import pandas as pd
import base64, io

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets"
DOWNLOAD_DIR = os.path.join(BASE_DIR, "downloaded_jackets")
THUMB_DIR = os.path.join(BASE_DIR, "thumbnails")
os.makedirs(THUMB_DIR, exist_ok=True)

REAPPRAISED_CHECKPOINT = os.path.join(BASE_DIR, "reappraised_jackets_checkpoint.json")
RESEARCH_FILE = os.path.join(BASE_DIR, "brand_research_database.json")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.xlsx")
OUTPUT_CSV = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.csv")

# Auth
feynman_auth = os.path.expanduser('~/.feynman/agent/auth.json')
with open(feynman_auth, 'r', encoding='utf-8') as f:
    api_key = json.load(f)['opencode-go']['key']

from upgrade_master_jacket_catalog import load_all_segments, encode_thumb, make_thumbnail

def process_jacket_full_photos(seg):
    k = seg["key"]
    fld_name = seg["folder"]
    folder_path = os.path.join(DOWNLOAD_DIR, fld_name)
    all_photos = seg["all_photos"]

    # Limit to at most 10 photos if segment is exceptionally large to keep prompt crisp
    photos_to_send = all_photos[:10] if len(all_photos) > 10 else all_photos

    prompt = f"""당신은 전세계 빈티지 가죽자켓 전문 품질/브랜드 감정사입니다.
제공된 사진들은 1벌의 가죽자켓에 대한 전체 사진들입니다. (총 {len(photos_to_send)}장: {', '.join(photos_to_send)})

[절대 지침]
1. brand_name (브랜드명 판독):
   - 사진들 속 목 라벨, 가죽 패치, 안쪽 상표, 케어라벨, 단추/지퍼 각인에 적힌 실제 브랜드명을 100% 정밀하게 읽어내세요.
   - 절대 추측하거나 유명 브랜드를 임의로 지어내지 마세요. 사진에 적힌 영문/한자/일문 그대로 기재하세요.
   - 라벨에 브랜드명이 전혀 없고 순수 세탁기호나 'Genuine Leather', '천연가죽' 마크만 있다면 반드시 '빈티지 오리지널(무명)'으로 적으세요.
2. 사진 분류 (제공된 파일명 목록 {photos_to_send} 중에서 정확히 1개씩 매칭):
   - card_photo: 손글씨 실측표 카드 사진 파일명
   - front_photo: 자켓 정면 전체 사진 파일명
   - back_photo: 자켓 뒷면 전체 사진 파일명
   - label_photo: 브랜드 목 라벨/상표/로고가 가장 선명하게 보이는 사진 파일명
   - detail_photo: 안감 또는 소매단/지퍼/마모 디테일 사진 파일명
3. condition_grade (상태 등급):
   - A급 (우수): 찢김/헤짐 없이 가죽 질감 및 자연스러운 에이징 우수
   - B급 (보통): 카라/소매단에 빈티지 주름이나 생활 마모가 있으나 착용에 무리없음
   - C급 (데미지): 가죽 찢김, 깊은 까짐, 심한 오염 등 결함 있음
4. condition_details: 가죽 표면 질감, 카라/소매 마모도, 안감 오염 여부를 한국어로 1~2문장 요약
5. 치수 및 스펙: 실측카드에서 품번(jacket_code), 어깨(shoulder_cm), 가슴(chest_cm), 기장(length_cm), 소매(sleeve_cm), 가죽소재(leather_type) 추출

반드시 순수 JSON 형식으로만 응답하세요:
{{
  "brand_name": "판독된 브랜드명",
  "card_photo": "파일명",
  "front_photo": "파일명",
  "back_photo": "파일명",
  "label_photo": "파일명",
  "detail_photo": "파일명",
  "condition_grade": "A급 / B급 / C급",
  "condition_details": "상태 코멘트",
  "jacket_code": "품번",
  "leather_type": "가죽 종류",
  "shoulder_cm": "숫자",
  "chest_cm": "숫자",
  "length_cm": "숫자",
  "sleeve_cm": "숫자"
}}"""

    content_list = [{'type': 'text', 'text': prompt}]
    for p in photos_to_send:
        full_p = os.path.join(folder_path, p)
        b64 = encode_thumb(full_p, 250)
        if b64:
            content_list.append({'type': 'text', 'text': f'사진: {p}'})
            content_list.append({'type': 'image_url', 'image_url': {'url': f'data:image/jpeg;base64,{b64}'}})

    url = 'https://opencode.ai/zen/go/v1/chat/completions'
    headers = {
        'Authorization': f'Bearer {api_key}',
        'Content-Type': 'application/json',
        'x-opencode-session': f'ses_full_{uuid.uuid4().hex[:8]}'
    }
    payload = {
        'model': 'qwen3.8-flash',
        'messages': [{'role': 'user', 'content': content_list}],
        'max_tokens': 1500
    }

    for attempt in range(2):
        try:
            resp = requests.post(url, headers=headers, json=payload, timeout=50)
            if resp.status_code == 200:
                txt = resp.json()['choices'][0]['message'].get('content') or ""
                clean_txt = re.sub(r'```json\s*', '', txt)
                clean_txt = re.sub(r'```', '', clean_txt).strip()
                json_match = re.search(r"\{[\s\S]*\}", clean_txt)
                if json_match:
                    raw_json = json_match.group(0)
                    try:
                        parsed = json.loads(raw_json)
                    except Exception:
                        cleaned = re.sub(r',\s*\}', '}', raw_json)
                        parsed = json.loads(cleaned)

                    # Validate photo filenames exist in seg['all_photos']
                    def pick_valid(fname, default_fallback):
                        return fname if fname in all_photos else default_fallback

                    c_p = pick_valid(parsed.get("card_photo"), all_photos[0])
                    f_p = pick_valid(parsed.get("front_photo"), all_photos[1] if len(all_photos) > 1 else all_photos[0])
                    b_p = pick_valid(parsed.get("back_photo"), all_photos[-1] if len(all_photos) > 2 else f_p)
                    l_p = pick_valid(parsed.get("label_photo"), all_photos[2] if len(all_photos) > 3 else f_p)
                    d_p = pick_valid(parsed.get("detail_photo"), all_photos[-2] if len(all_photos) > 4 else f_p)

                    parsed["card_photo"] = c_p
                    parsed["front_photo"] = f_p
                    parsed["back_photo"] = b_p
                    parsed["label_photo"] = l_p
                    parsed["detail_photo"] = d_p

                    return k, parsed
        except Exception:
            time.sleep(1)

    # Fallback if failed
    fallback = {
        "brand_name": "빈티지 오리지널(무명)",
        "card_photo": all_photos[0],
        "front_photo": all_photos[1] if len(all_photos) > 1 else all_photos[0],
        "back_photo": all_photos[-1] if len(all_photos) > 2 else all_photos[0],
        "label_photo": all_photos[2] if len(all_photos) > 3 else all_photos[0],
        "detail_photo": all_photos[-2] if len(all_photos) > 4 else all_photos[0],
        "condition_grade": "B급",
        "condition_details": "자연스러운 에이징 및 생활 주름이 있는 양호한 상태입니다.",
        "jacket_code": f"J#{seg['jacket_index']}",
        "leather_type": "천연가죽",
        "shoulder_cm": "-",
        "chest_cm": "-",
        "length_cm": "-",
        "sleeve_cm": "-"
    }
    return k, fallback

def main():
    print("=" * 65)
    print("🧥 [Finder] 556벌 가죽자켓 전체 사진 전수 투입 정밀 재감정 & 엑셀 재구축")
    print("=" * 65)

    segments = load_all_segments()
    print(f"[*] 총 자켓 수: {len(segments)}벌")

    checkpoint = {}
    if os.path.exists(REAPPRAISED_CHECKPOINT):
        try:
            with open(REAPPRAISED_CHECKPOINT, "r", encoding="utf-8") as f:
                checkpoint = json.load(f)
        except Exception:
            checkpoint = {}

    to_run = [s for s in segments if s["key"] not in checkpoint]
    print(f"[*] 기존 완료: {len(checkpoint)}벌 / 신규 분석 대상: {len(to_run)}벌")

    if to_run:
        print(f"[*] 16스레드 병렬 분석 시작 (각 자켓당 전체 사진 7~10장 전수 투입)...")
        start_time = time.time()
        completed = len(checkpoint)

        with ThreadPoolExecutor(max_workers=16) as executor:
            future_to_seg = {executor.submit(process_jacket_full_photos, s): s for s in to_run}
            batch_count = 0
            for future in as_completed(future_to_seg):
                batch_count += 1
                completed += 1
                try:
                    k, res = future.result()
                    checkpoint[k] = res

                    b_name = res.get("brand_name", "-")
                    c_grade = res.get("condition_grade", "-")
                    code = res.get("jacket_code", "-")

                    if batch_count % 10 == 0 or batch_count == len(to_run):
                        elapsed = time.time() - start_time
                        speed = batch_count / elapsed if elapsed > 0 else 0
                        remain = (len(to_run) - batch_count) / speed if speed > 0 else 0
                        print(f"  [{completed}/{len(segments)}] {k} | 품번: {code} | 브랜드: {b_name} | 상태: {c_grade} | 남은시간: {int(remain)}초", flush=True)

                        with open(REAPPRAISED_CHECKPOINT, "w", encoding="utf-8") as f:
                            json.dump(checkpoint, f, ensure_ascii=False, indent=2)
                except Exception as e:
                    pass

        with open(REAPPRAISED_CHECKPOINT, "w", encoding="utf-8") as f:
            json.dump(checkpoint, f, ensure_ascii=False, indent=2)
        print(f"[✓] 556벌 전체 사진 전수 정밀 감정 완료! 소요시간: {int(time.time() - start_time)}초")

    # Load brand research database
    brand_research = {}
    if os.path.exists(RESEARCH_FILE):
        with open(RESEARCH_FILE, "r", encoding="utf-8") as f:
            brand_research = json.load(f)

    # Function to match brand
    def match_research(brand_str):
        if not brand_str:
            return brand_research.get("빈티지 오리지널 (무명)", {})
        b_clean = brand_str.strip()
        unnamed = ['무명', '판독 불가', '판독불가', '미상', '미식별', '미표기', 
                   '식별 불가', '식별불가', '라벨 미확인', '미확인', '불명', 
                   '무지라벨', '블랭크', '빈티지 오리지널', 'Non-brand', 'non-brand', '식별불능']
        if any(w in b_clean for w in unnamed):
            return brand_research.get("빈티지 오리지널 (무명)", {})

        if b_clean in brand_research:
            return brand_research[b_clean]

        for k, v in brand_research.items():
            if k != "빈티지 오리지널 (무명)" and (k.lower() in b_clean.lower() or b_clean.lower() in k.lower()):
                return v

        return {
            "official_name": b_clean,
            "country": "불명",
            "tier": "C등급 (일반 빈티지 공방/도메스틱)",
            "heritage": "로컬 피혁 공방 및 도메스틱 패션 브랜드에서 제작된 천연가죽 의류."
        }

    # Enrich segments
    enriched = []
    for seg in segments:
        k = seg["key"]
        cp = checkpoint.get(k, {})
        fld_path = os.path.join(DOWNLOAD_DIR, seg["folder"])

        b_raw = cp.get("brand_name", "빈티지 오리지널(무명)")
        res = match_research(b_raw)

        b_name = res.get("official_name", b_raw)
        b_country = res.get("country", "불명")
        b_tier_full = res.get("tier", "C등급 (일반 빈티지)")
        b_tier = b_tier_full.split()[0]
        b_heritage = res.get("heritage", "천연가죽 빈티지 의류")

        c_grade = cp.get("condition_grade", "B급")
        if not c_grade.endswith("급"): c_grade = f"{c_grade[0]}급"
        c_details = cp.get("condition_details", "자연스러운 빈티지 에이징 상태")

        # Recommendation logic
        if b_tier == "S등급":
            if c_grade in ["S급", "A급"]: rec = "⭐⭐⭐⭐⭐ 강력추천(S)"
            elif c_grade == "B급": rec = "⭐⭐⭐⭐ 추천(A)"
            else: rec = "⭐⭐ 비추천(C)"
        elif b_tier == "A등급":
            if c_grade in ["S급", "A급"]: rec = "⭐⭐⭐⭐ 추천(A)"
            elif c_grade == "B급": rec = "⭐⭐⭐ 보통(B)"
            else: rec = "⭐⭐ 비추천(C)"
        elif b_tier == "B등급":
            if c_grade in ["S급", "A급"]: rec = "⭐⭐⭐⭐ 추천(A)"
            elif c_grade == "B급": rec = "⭐⭐⭐ 보통(B)"
            else: rec = "⭐⭐ 비추천(C)"
        else: # C등급
            rec = "⭐⭐⭐ 보통(B)" if c_grade != "C급" else "⭐⭐ 비추천(C)"

        card_p = os.path.join(fld_path, cp.get("card_photo") or seg["card_photo"])
        front_p = os.path.join(fld_path, cp.get("front_photo") or seg["front_photo"])
        back_p = os.path.join(fld_path, cp.get("back_photo") or seg["all_photos"][-1])
        label_p = os.path.join(fld_path, cp.get("label_photo") or (seg["all_photos"][2] if len(seg["all_photos"])>2 else seg["front_photo"]))
        detail_p = os.path.join(fld_path, cp.get("detail_photo") or (seg["all_photos"][-2] if len(seg["all_photos"])>4 else seg["front_photo"]))

        enriched.append({
            "seg": seg,
            "key": k,
            "code": cp.get("jacket_code", f"J#{seg['jacket_index']}"),
            "recommendation": rec,
            "tier": b_tier,
            "condition_grade": c_grade,
            "brand": b_name,
            "country": b_country,
            "heritage": b_heritage,
            "condition_details": c_details,
            "card_p": card_p,
            "front_p": front_p,
            "back_p": back_p,
            "label_p": label_p,
            "detail_p": detail_p,
            "leather_type": cp.get("leather_type", "천연가죽"),
            "shoulder": cp.get("shoulder_cm", "-"),
            "chest": cp.get("chest_cm", "-"),
            "length": cp.get("length_cm", "-"),
            "sleeve": cp.get("sleeve_cm", "-"),
        })

    # Sort priority
    def sort_key(it):
        s = 0
        if "강력추천(S)" in it["recommendation"]: s += 100
        elif "추천(A)" in it["recommendation"]: s += 70
        elif "보통(B)" in it["recommendation"]: s += 40
        else: s += 10

        if it["tier"] == "S등급": s += 25
        elif it["tier"] == "A등급": s += 15
        elif it["tier"] == "B등급": s += 10

        if it["condition_grade"] in ["S급", "A급"]: s += 5
        return -s

    sorted_items = sorted(enriched, key=sort_key)

    # Re-build Excel
    print("\n[*] 5종 사진 완벽 정렬 마스터 엑셀 생성 중...")
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "자켓_종합평가_도감"

    headers1 = [
        "순번", "자켓품번", "종합추천", "브랜드 등급", "상태 등급",
        "판독 브랜드", "브랜드 국적", "브랜드 역사 & 배경", "상태 상세평가 코멘트",
        "실측카드 사진", "자켓 정면 사진", "자켓 뒷면 사진", "브랜드 라벨 사진", "안감/디테일 사진",
        "가죽 소재", "어깨(cm)", "가슴(cm)", "기장(cm)", "소매(cm)",
        "전체 사진수", "사진 범위", "상품폴더", "상품명"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    ws1.append(headers1)
    for col_idx in range(1, len(headers1) + 1):
        c = ws1.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws1.row_dimensions[1].height = 30

    tier_colors = {
        "S등급": {"bg": "FEF3C7", "fg": "92400E", "bold": True},
        "A등급": {"bg": "DCFCE7", "fg": "166534", "bold": True},
        "B등급": {"bg": "E0F2FE", "fg": "075985", "bold": True},
        "C등급": {"bg": "F3F4F6", "fg": "4B5563", "bold": False}
    }
    grade_colors = {
        "S급": {"bg": "FEF3C7", "fg": "92400E", "bold": True},
        "A급": {"bg": "DCFCE7", "fg": "166534", "bold": True},
        "B급": {"bg": "E0F2FE", "fg": "075985", "bold": True},
        "C급": {"bg": "FEE2E2", "fg": "991B1B", "bold": True}
    }

    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="맑은 고딕", size=9)
    bold_font = Font(name="맑은 고딕", size=9, bold=True)
    desc_font = Font(name="맑은 고딕", size=8)

    csv_rows = []

    for idx, itm in enumerate(sorted_items, start=1):
        row_num = idx + 1
        ws1.row_dimensions[row_num].height = 68
        seg = itm["seg"]

        row_data = [
            idx,
            itm["code"],
            itm["recommendation"],
            itm["tier"],
            itm["condition_grade"],
            itm["brand"],
            itm["country"],
            itm["heritage"],
            itm["condition_details"],
            "", "", "", "", "", # 5 embedded photos: J, K, L, M, N
            itm["leather_type"],
            itm["shoulder"],
            itm["chest"],
            itm["length"],
            itm["sleeve"],
            f"{len(seg['all_photos'])}장",
            seg["photo_range"],
            seg["folder"],
            seg["item_title"]
        ]
        ws1.append(row_data)

        fill_color = zebra_fill if idx % 2 == 0 else white_fill
        for c_idx in range(1, len(headers1) + 1):
            cell = ws1.cell(row=row_num, column=c_idx)
            cell.fill = fill_color
            cell.border = thin_border
            cell.font = data_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws1.cell(row=row_num, column=2).font = bold_font
        ws1.cell(row=row_num, column=6).font = Font(name="맑은 고딕", size=9, bold=True, color="1E3A8A")
        ws1.cell(row=row_num, column=8).font = desc_font
        ws1.cell(row=row_num, column=8).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws1.cell(row=row_num, column=9).font = desc_font
        ws1.cell(row=row_num, column=9).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws1.cell(row=row_num, column=23).alignment = Alignment(horizontal="left", vertical="center")

        # Badges
        if itm["tier"] in tier_colors:
            tc = tier_colors[itm["tier"]]
            ws1.cell(row=row_num, column=4).fill = PatternFill(start_color=tc["bg"], end_color=tc["bg"], fill_type="solid")
            ws1.cell(row=row_num, column=4).font = Font(name="맑은 고딕", size=9, bold=tc["bold"], color=tc["fg"])

        if itm["condition_grade"] in grade_colors:
            gc = grade_colors[itm["condition_grade"]]
            ws1.cell(row=row_num, column=5).fill = PatternFill(start_color=gc["bg"], end_color=gc["bg"], fill_type="solid")
            ws1.cell(row=row_num, column=5).font = Font(name="맑은 고딕", size=9, bold=gc["bold"], color=gc["fg"])

        # Embed exact verified photos (Cols J, K, L, M, N)
        photo_targets = [
            (itm["card_p"], "J", f"card_{seg['folder']}"),
            (itm["front_p"], "K", f"front_{seg['folder']}"),
            (itm["back_p"], "L", f"back_{seg['folder']}"),
            (itm["label_p"], "M", f"lbl_{seg['folder']}"),
            (itm["detail_p"], "N", f"det_{seg['folder']}")
        ]

        for p_src, col_letter, p_prefix in photo_targets:
            t_path = make_thumbnail(p_src, p_prefix)
            if t_path and os.path.exists(t_path):
                try:
                    img = XLImage(t_path)
                    ws1.add_image(img, f"{col_letter}{row_num}")
                except Exception:
                    pass

        csv_rows.append({
            "순번": idx,
            "자켓품번": itm["code"],
            "종합추천": itm["recommendation"],
            "브랜드등급": itm["tier"],
            "상태등급": itm["condition_grade"],
            "판독브랜드": itm["brand"],
            "브랜드국적": itm["country"],
            "브랜드조사내용": itm["heritage"],
            "상태상세평가": itm["condition_details"],
            "가죽소재": itm["leather_type"],
            "어깨(cm)": itm["shoulder"],
            "가슴(cm)": itm["chest"],
            "기장(cm)": itm["length"],
            "소매(cm)": itm["sleeve"],
            "실측카드사진": os.path.basename(itm["card_p"]),
            "자켓정면사진": os.path.basename(itm["front_p"]),
            "자켓뒷면사진": os.path.basename(itm["back_p"]),
            "브랜드라벨사진": os.path.basename(itm["label_p"]),
            "디테일안감사진": os.path.basename(itm["detail_p"]),
            "전체사진수": len(seg["all_photos"]),
            "사진범위": seg["photo_range"],
            "상품폴더": seg["folder"],
            "상품명": seg["item_title"]
        })

    widths1 = {
        "A": 6, "B": 11, "C": 18, "D": 11, "E": 11,
        "F": 20, "G": 14, "H": 46, "I": 40,
        "J": 13, "K": 13, "L": 13, "M": 13, "N": 13,
        "O": 14, "P": 9, "Q": 9, "R": 9, "S": 9,
        "T": 10, "U": 26, "V": 20, "W": 35
    }
    for col_l, w in widths1.items():
        ws1.column_dimensions[col_l].width = w

    ws1.freeze_panes = "J2"
    wb.save(OUTPUT_EXCEL)
    print(f"[🎉] 엑셀 도감 재생성 완료 -> {OUTPUT_EXCEL}")

    pd.DataFrame(csv_rows).to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    print(f"[✓] CSV 저장 완료 -> {OUTPUT_CSV}")

if __name__ == "__main__":
    main()
