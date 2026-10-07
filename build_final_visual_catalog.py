import os, glob, sys, json, requests, uuid, re
from concurrent.futures import ThreadPoolExecutor, as_completed
from PIL import Image
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
import pandas as pd
import base64, io

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = r"c:\Users\passp\Desktop\univercity\4-2\taobao_jackets"
DOWNLOAD_DIR = os.path.join(BASE_DIR, "downloaded_jackets")
THUMB_DIR = os.path.join(BASE_DIR, "thumbnails")
os.makedirs(THUMB_DIR, exist_ok=True)

CACHE_CARDS_FILE = os.path.join(BASE_DIR, "detected_cards_cache_v2.json")
CHECKPOINT_FILE = os.path.join(BASE_DIR, "analyzed_jackets_checkpoint.json")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.xlsx")
OUTPUT_CSV = os.path.join(BASE_DIR, "taobao_leather_jackets_with_brands.csv")

feynman_auth = os.path.expanduser('~/.feynman/agent/auth.json')
with open(feynman_auth, 'r', encoding='utf-8') as f:
    api_key = json.load(f)['opencode-go']['key']

def encode_image(img_path, max_dim=450):
    try:
        with Image.open(img_path) as im:
            im = im.convert("RGB")
            w, h = im.size
            if max(w, h) > max_dim:
                scale = max_dim / float(max(w, h))
                im = im.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
            buf = io.BytesIO()
            im.save(buf, format="JPEG", quality=80)
            return base64.b64encode(buf.getvalue()).decode("utf-8")
    except Exception:
        return None

def make_thumbnail(src_path, prefix, max_size=(75, 75)):
    if not src_path or not os.path.exists(src_path):
        return None
    fname = f"{prefix}_{os.path.basename(src_path)}"
    out_path = os.path.join(THUMB_DIR, fname)
    if os.path.exists(out_path):
        return out_path
    try:
        with Image.open(src_path) as im:
            im = im.convert("RGB")
            im.thumbnail(max_size, Image.Resampling.LANCZOS)
            im.save(out_path, "JPEG", quality=85)
        return out_path
    except Exception:
        return None

def analyze_jacket_worker(seg):
    folder_path = os.path.join(DOWNLOAD_DIR, seg["folder"])
    card_path = os.path.join(folder_path, seg["card_photo"])
    front_path = os.path.join(folder_path, seg["front_photo"]) if seg.get("front_photo") else None
    label_path = os.path.join(folder_path, seg["label_photo"]) if seg.get("label_photo") else None

    # Check if thumbnail already made
    thumb_card = make_thumbnail(card_path, f"{seg['folder']}_card")
    thumb_front = make_thumbnail(front_path, f"{seg['folder']}_front") if front_path else None

    card_b64 = encode_image(card_path)
    front_b64 = encode_image(front_path) if front_path and os.path.exists(front_path) else None
    label_b64 = encode_image(label_path) if label_path and os.path.exists(label_path) else None

    prompt = (
        "제공된 사진들은 1벌의 빈티지 가죽자켓(실측표 카드, 자켓 정면, 라벨)입니다.\n"
        "장황한 설명 없이 아래 JSON 포맷으로만 응답해주세요:\n"
        "```json\n"
        "{\n"
        '  "jacket_code": "실측지 상단 编号 (예: P211, AC91, E1 등. 없으면 미상)",\n'
        '  "brand": "라벨에 적힌 정확한 브랜드명 영문/한글 (미상일 경우 빈티지 오리지널)",\n'
        '  "leather_type": "가죽 종류 (양가죽, 소가죽, 사슴가죽, 돈모 등)",\n'
        '  "origin": "원산지 (미국, 이탈리아, 일본, 한국, 중국 등, 불명 시 불명)",\n'
        '  "shoulder_cm": "어깨 실측 숫자 (단위 제외)",\n'
        '  "chest_cm": "가슴 실측 숫자 (단위 제외)",\n'
        '  "length_cm": "총기장 실측 숫자 (단위 제외)",\n'
        '  "sleeve_cm": "소매길이 실측 숫자 (단위 제외)"\n'
        "}\n"
        "```"
    )

    content_list = [{"type": "text", "text": prompt}]
    if card_b64:
        content_list.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{card_b64}"}})
    if front_b64:
        content_list.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{front_b64}"}})
    if label_b64:
        content_list.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{label_b64}"}})

    url = "https://opencode.ai/zen/go/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "x-opencode-session": f"ses_jkt_{uuid.uuid4().hex[:8]}"
    }
    payload = {
        "model": "deepseek-v4.1-flash",
        "messages": [{"role": "user", "content": content_list}],
        "max_tokens": 1000
    }

    result_data = {
        "jacket_code": "",
        "brand": "빈티지 오리지널",
        "leather_type": "천연가죽",
        "origin": "불명",
        "shoulder_cm": "-",
        "chest_cm": "-",
        "length_cm": "-",
        "sleeve_cm": "-"
    }

    try:
        resp = requests.post(url, headers=headers, json=payload, timeout=40)
        if resp.status_code == 200:
            c = resp.json()["choices"][0]["message"]
            txt = c.get("content") or c.get("reasoning_content") or ""
            m = re.search(r'```json\s*(\{.*?\})\s*```', txt, re.DOTALL)
            if not m:
                m = re.search(r'(\{.*?"jacket_code".*?\})', txt, re.DOTALL)
            if not m:
                m = re.search(r'(\{.*?"brand".*?\})', txt, re.DOTALL)
            if m:
                parsed = json.loads(m.group(1))
                for k in result_data:
                    if k in parsed and parsed[k]:
                        val = str(parsed[k]).strip()
                        if val.lower() not in ["null", "none", "미상", "불명", ""]:
                            result_data[k] = val
    except Exception as e:
        print(f"Error on {seg['folder']} {seg['card_photo']}: {e}")

    seg_copy = dict(seg)
    seg_copy.update(result_data)
    seg_copy["thumb_card"] = thumb_card
    seg_copy["thumb_front"] = thumb_front
    return seg_copy

def load_all_segments():
    with open(CACHE_CARDS_FILE, "r", encoding="utf-8") as f:
        cache_data = json.load(f)

    folders = sorted(glob.glob(os.path.join(DOWNLOAD_DIR, "*")))
    all_segs = []

    for fld in folders:
        fld_name = os.path.basename(fld)
        jpgs = sorted(glob.glob(os.path.join(fld, "photo_*.jpg")))
        if not jpgs or fld_name not in cache_data:
            continue

        item_title = fld_name
        info_file = os.path.join(fld, "item_info.txt")
        if os.path.exists(info_file):
            with open(info_file, "r", encoding="utf-8", errors="ignore") as inf:
                for l in inf:
                    if l.startswith("상품명:"): item_title = l.replace("상품명:", "").strip()

        cards = cache_data[fld_name]
        # Filter guide banners and adjacent duplicates
        valid_cards = []
        for c in cards:
            c_path = os.path.join(fld, c)
            if os.path.exists(c_path) and os.path.getsize(c_path) != 33084:
                valid_cards.append(c)

        # Filter adjacent
        filtered_cards = []
        for c in valid_cards:
            num = int(c.replace("photo_", "").replace(".jpg", ""))
            if filtered_cards:
                prev_num = int(filtered_cards[-1].replace("photo_", "").replace(".jpg", ""))
                if num - prev_num < 3:
                    continue
            filtered_cards.append(c)

        fname_list = [os.path.basename(p) for p in jpgs]

        for idx, card_fname in enumerate(filtered_cards):
            card_idx = fname_list.index(card_fname) if card_fname in fname_list else -1
            if card_idx == -1: continue

            if idx + 1 < len(filtered_cards):
                next_card_idx = fname_list.index(filtered_cards[idx + 1]) if filtered_cards[idx + 1] in fname_list else len(fname_list)
            else:
                next_card_idx = len(fname_list)

            jkt_photos = fname_list[card_idx:next_card_idx]
            front_photo = jkt_photos[1] if len(jkt_photos) > 1 else card_fname
            label_photo = jkt_photos[-1] if len(jkt_photos) > 2 else front_photo
            for candidate_label in jkt_photos[2:]:
                label_photo = candidate_label
                break

            seg = {
                "folder": fld_name,
                "jacket_index": idx + 1,
                "total_in_folder": len(filtered_cards),
                "card_photo": card_fname,
                "front_photo": front_photo,
                "label_photo": label_photo,
                "photo_range": f"{jkt_photos[0]} ~ {jkt_photos[-1]} ({len(jkt_photos)}장)",
                "item_title": item_title
            }
            all_segs.append(seg)

    return all_segs

def main():
    segments = load_all_segments()
    print(f"[*] Total segmented jackets loaded: {len(segments)}")

    # Load existing checkpoint
    completed_data = {}
    if os.path.exists(CHECKPOINT_FILE):
        try:
            with open(CHECKPOINT_FILE, "r", encoding="utf-8") as cf:
                completed_data = json.load(cf)
            print(f"[*] Loaded {len(completed_data)} already analyzed jackets from checkpoint.")
        except Exception:
            completed_data = {}

    to_analyze = []
    for s in segments:
        key = f"{s['folder']}_{s['card_photo']}"
        if key not in completed_data:
            to_analyze.append(s)

    print(f"[*] Jackets remaining to analyze with OpenCode Go: {len(to_analyze)}")

    if to_analyze:
        # Run concurrent analysis with 8 threads
        with ThreadPoolExecutor(max_workers=8) as executor:
            future_to_seg = {executor.submit(analyze_jacket_worker, seg): seg for seg in to_analyze}
            count = 0
            for future in as_completed(future_to_seg):
                count += 1
                try:
                    res = future.result()
                    key = f"{res['folder']}_{res['card_photo']}"
                    completed_data[key] = res
                    if count % 10 == 0 or count == len(to_analyze):
                        print(f"  [{count}/{len(to_analyze)}] Analyzed: {res['jacket_code']} | Brand: {res['brand']} | Mat: {res['leather_type']}", flush=True)
                        with open(CHECKPOINT_FILE, "w", encoding="utf-8") as cf:
                            json.dump(completed_data, cf, ensure_ascii=False, indent=2)
                except Exception as e:
                    print(f"Worker exception: {e}")

        # Final checkpoint save
        with open(CHECKPOINT_FILE, "w", encoding="utf-8") as cf:
            json.dump(completed_data, cf, ensure_ascii=False, indent=2)

    print(f"[+] All {len(segments)} jackets analyzed! Now generating Excel with embedded photos...")

    # Build final Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "빈티지_가죽자켓_실물도감"

    headers = [
        "순번", "자켓품번", "실측카드 사진", "자켓 정면 사진", 
        "판독브랜드", "가죽소재", "원산지", 
        "어깨(cm)", "가슴(cm)", "기장(cm)", "소매(cm)", 
        "사진범위", "상품폴더", "상품명"
    ]

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 30

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    col_widths = {
        1: 8,   # 순번
        2: 14,  # 자켓품번
        3: 15,  # 실측카드 사진
        4: 15,  # 자켓 정면 사진
        5: 20,  # 판독브랜드
        6: 18,  # 가죽소재
        7: 15,  # 원산지
        8: 12,  # 어깨
        9: 12,  # 가슴
        10: 12, # 기장
        11: 12, # 소매
        12: 24, # 사진범위
        13: 22, # 상품폴더
        14: 38  # 상품명
    }
    for col_idx, width in col_widths.items():
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    csv_rows = []

    for idx, seg in enumerate(segments, start=1):
        key = f"{seg['folder']}_{seg['card_photo']}"
        data = completed_data.get(key, seg)

        code = data.get("jacket_code") or f"{seg['folder'].replace('item_', '')}-{seg['jacket_index']}"
        brand = data.get("brand") or "빈티지 오리지널"
        leather = data.get("leather_type") or "천연가죽"
        origin = data.get("origin") or "불명"
        sh = data.get("shoulder_cm") or "-"
        ch = data.get("chest_cm") or "-"
        ln = data.get("length_cm") or "-"
        sl = data.get("sleeve_cm") or "-"

        row_num = idx + 1
        ws.row_dimensions[row_num].height = 75

        row_vals = [
            idx, code, "", "", # C and D for embedded images
            brand, leather, origin,
            sh, ch, ln, sl,
            seg["photo_range"], seg["folder"], seg["item_title"]
        ]
        ws.append(row_vals)

        for col_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = Font(name="맑은 고딕", size=10)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Column C: Embed Size Card Image
        thumb_card = data.get("thumb_card")
        if not thumb_card or not os.path.exists(thumb_card):
            thumb_card = make_thumbnail(os.path.join(DOWNLOAD_DIR, seg["folder"], seg["card_photo"]), f"{seg['folder']}_card")
        if thumb_card and os.path.exists(thumb_card):
            xl_c = XLImage(thumb_card)
            xl_c.width = 70
            xl_c.height = 70
            ws.add_image(xl_c, f"C{row_num}")

        # Column D: Embed Jacket Front Image
        thumb_front = data.get("thumb_front")
        if not thumb_front or not os.path.exists(thumb_front):
            thumb_front = make_thumbnail(os.path.join(DOWNLOAD_DIR, seg["folder"], seg["front_photo"]), f"{seg['folder']}_front")
        if thumb_front and os.path.exists(thumb_front):
            xl_f = XLImage(thumb_front)
            xl_f.width = 70
            xl_f.height = 70
            ws.add_image(xl_f, f"D{row_num}")

        csv_rows.append({
            "순번": idx,
            "자켓품번": code,
            "실측카드사진": seg["card_photo"],
            "자켓정면사진": seg["front_photo"],
            "판독브랜드": brand,
            "가죽소재": leather,
            "원산지": origin,
            "어깨(cm)": sh,
            "가슴(cm)": ch,
            "기장(cm)": ln,
            "소매(cm)": sl,
            "사진범위": seg["photo_range"],
            "상품폴더": seg["folder"],
            "상품명": seg["item_title"]
        })

    wb.save(OUTPUT_EXCEL)
    print(f"[🎉] Visual Excel created successfully at:\n  - {OUTPUT_EXCEL}")

    df_csv = pd.DataFrame(csv_rows)
    df_csv.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    print(f"[🎉] CSV created successfully at:\n  - {OUTPUT_CSV}")

if __name__ == "__main__":
    main()
